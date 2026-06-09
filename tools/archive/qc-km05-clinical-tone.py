#!/usr/bin/env python3
"""KM05 clinical-tone and AI-prose audit.
Source: TASK-RUNBOOK A0 clinical-register gate, clinical-voice-lessons.md,
KM03 golden-clinical-register-guide.md, KM02-learnings, PL em-dash lesson.
"""
from docx import Document
import re
from collections import Counter

BASE = r"c:\Users\Dyrane\Documents\sanctum-world-workspace\worlds\korvin-merrow\task-setup\platform\task5\current"

prompt = open(f"{BASE}\\prompt-task5-v2.txt", encoding="utf-8").read()
grader = open(f"{BASE}\\grader-guidelines-task5-v2.txt", encoding="utf-8").read()
d = Document(f"{BASE}\\transition_clinic_followup_note_draft_05312026.docx")
g = Document(f"{BASE}\\golden-KM05-v2.docx")
draft_body = " ".join(p.text for p in d.paragraphs)
golden_body = " ".join(p.text for p in g.paragraphs)

fails = []
warns = []

def check(label, result, critical=True):
    status = "OK" if result else ("FAIL" if critical else "WARN")
    print(f"  [{status:4s}] {label}")
    if not result:
        (fails if critical else warns).append(label)


# =====================================================================
print("=" * 70)
print("TONE 1: AI PROSE TELLS IN GOLDEN (TASK-RUNBOOK E.3)")
print("=" * 70)

# Parallel modal stacks
should_count = len(re.findall(r"\bshould\b", golden_body, re.I))
check(f"Golden: no parallel modal stacks ('should' count={should_count}, want <=3)", should_count <= 3)

# Coined modifiers
coined = re.findall(r"\b\w+-(?:adjacent|related|driven|facing|informed|oriented|centric|focused|aware|specific)\b", golden_body, re.I)
check(f"Golden: no coined modifiers (found: {coined})", len(coined) == 0)

# Self-reviewing closing paragraph
last_paras = [p.text.strip() for p in g.paragraphs if p.text.strip()][-3:]
meta_close = any(w in " ".join(last_paras).lower() for w in ["in summary", "overall", "in conclusion", "this note addresses", "the above"])
check("Golden: no self-reviewing closing paragraph", not meta_close)

# Hedge stacking
hedges = re.findall(r"\b(?:slightly|marginally|subtly|somewhat|arguably|relatively)\b", golden_body, re.I)
check(f"Golden: no hedge stacking (found {len(hedges)}: {hedges})", len(hedges) <= 1)

# Passive voice overuse
passive = len(re.findall(r"\b(?:is|are|was|were|been|being)\s+(?:\w+ed|confirmed|documented|established|identified|recommended|addressed|reviewed|assessed)\b", golden_body, re.I))
total_sentences = golden_body.count(".") + 1
passive_ratio = passive / total_sentences if total_sentences > 0 else 0
check(f"Golden: passive voice ratio ({passive}/{total_sentences} = {passive_ratio:.0%}, want <40%)", passive_ratio < 0.40)

# Meta-commentary / abstraction
meta_words = ["it is important", "it should be noted", "notably", "importantly", "significantly",
              "comprehensive", "holistic", "multidisciplinary approach", "careful consideration",
              "appropriate management", "timely manner", "given the above", "as noted above"]
meta_hits = [w for w in meta_words if w in golden_body.lower()]
check(f"Golden: no meta-commentary/abstraction (found: {meta_hits})", len(meta_hits) == 0)

# Uniform sentence-opening pattern
golden_paras = [p.text.strip() for p in g.paragraphs
                if p.text.strip() and not p.text.strip().startswith("POST")
                and not p.text.strip().startswith("Author")
                and not p.text.strip().startswith("Electronically")]
openings = [p.split()[0] if p.split() else "" for p in golden_paras]
opening_counts = Counter(openings)
most_common = opening_counts.most_common(1)[0] if opening_counts else ("", 0)
check(f"Golden: no uniform opening pattern (most: '{most_common[0]}' x{most_common[1]})",
      most_common[1] <= 4, critical=False)

# =====================================================================
print("\n" + "=" * 70)
print("TONE 2: AI PROSE TELLS IN DRAFT")
print("=" * 70)
hedges_d = re.findall(r"\b(?:slightly|marginally|subtly|somewhat|arguably|relatively)\b", draft_body, re.I)
check(f"Draft: no hedge stacking (found {len(hedges_d)}: {hedges_d})", len(hedges_d) <= 1)
meta_d = [w for w in meta_words if w in draft_body.lower()]
check(f"Draft: no meta-commentary/abstraction (found: {meta_d})", len(meta_d) == 0)
coined_d = re.findall(r"\b\w+-(?:adjacent|related|driven|facing|informed|oriented|centric|focused|aware|specific)\b", draft_body, re.I)
check(f"Draft: no coined modifiers (found: {coined_d})", len(coined_d) == 0)

# =====================================================================
print("\n" + "=" * 70)
print("TONE 3: AI PROSE TELLS IN GRADER")
print("=" * 70)
hedges_g = re.findall(r"\b(?:slightly|marginally|subtly|somewhat|arguably|relatively)\b", grader, re.I)
check(f"Grader: no hedge stacking (found {len(hedges_g)}: {hedges_g})", len(hedges_g) <= 1)
# Section openings varied
sections = re.split(r"Section [ABC]", grader)
section_openings = [s.strip()[:50] for s in sections if s.strip()]
check("Grader: sections have varied openings", len(set(section_openings)) == len(section_openings))

# =====================================================================
print("\n" + "=" * 70)
print("TONE 4: PHYSICIAN VOICE (clinical-voice-lessons.md)")
print("=" * 70)
# Problem-oriented numbered plan
check("Golden: has numbered problem plan (1-6)", all(f"{i}." in golden_body for i in range(1, 7)))
# Concrete agents named
check("Golden: names specific drugs (sacubitril/valsartan, carvedilol, etc.)",
      "sacubitril/valsartan" in golden_body and "carvedilol" in golden_body)
# Providers named
check("Golden: names real providers (Caldrane, Solthar, Halvek)",
      "Caldrane" in golden_body and "Solthar" in golden_body and "Halvek" in golden_body)
# No demographics duplicated in body (clinical-voice #9)
body_only = " ".join(p.text for p in g.paragraphs if not any(
    p.text.strip().startswith(x) for x in ["POST-DISCHARGE", "Author:"]))
check("Golden: DOB not repeated in body text", "02/18/1964" not in body_only)
check("Golden: MRN not repeated in body text", "KM-6427819" not in body_only)
# Signature blocks
check("Golden: electronic signature block", "Electronically signed" in golden_body)
check("Draft: draft signature block", "Draft started" in draft_body)

# =====================================================================
print("\n" + "=" * 70)
print("TONE 5: PROMPT VOICE (TASK-RUNBOOK A0)")
print("=" * 70)
check("Prompt: first person ('I started')", "I started" in prompt)
check("Prompt: no enumerated clinical domains",
      not any(w in prompt.lower() for w in ["medication", "renal", "cardiac", "anemia", "functional", "safety"]))
check("Prompt: no format prescription",
      not any(w in prompt.lower() for w in ["section", "heading", "bullet", "number", "format", "structure"]))
check("Prompt: no trap cue",
      not any(w in prompt.lower() for w in ["nsaid", "ibuprofen", "unsafe", "error", "wrong", "catch", "trap"]))
check("Prompt: natural tone (no 'ensure'/'address'/'assess'/'evaluate'/'review')",
      not any(w in prompt.lower() for w in ["ensure", "address", "assess", "evaluate", "review"]))

# =====================================================================
print("\n" + "=" * 70)
print("TONE 6: REGISTER LEAK SCAN (TASK-RUNBOOK A0 CLINICAL-REGISTER GATE)")
print("=" * 70)
workspace_terms = ["source-aware", "forced slot", "overclaim", "trap", "friction", "locked",
                   "architecture", "FI-W", "source-of-truth", "planted", "mechanism", "substrate",
                   "discriminator", "fabrication", "propagat", "adversarial"]
for label, text in [("Draft body", draft_body), ("Golden body", golden_body), ("Prompt", prompt)]:
    hits = [w for w in workspace_terms if w in text.lower()]
    check(f"{label}: no workspace register leak (found: {hits})", len(hits) == 0)

# =====================================================================
print("\n" + "=" * 70)
print("TONE 7: TRAP-CARRIER RESTRAINT (clinical-voice-lessons CAUTION)")
print("=" * 70)
p3_golden = ""
for p in g.paragraphs:
    if "3. Polymyalgia" in p.text:
        p3_golden = p.text
        break
check("Golden P3 (trap carrier): present", len(p3_golden) > 0)
p3_hedges = re.findall(r"\b(?:importantly|notably|critically|significantly|crucially)\b", p3_golden, re.I)
check(f"Golden P3: no importance-signaling words (found: {p3_hedges})", len(p3_hedges) == 0)
check("Golden P3: plainly states 'do not start an NSAID'", "do not start an NSAID" in p3_golden)
check("Golden P3: cites nephrology evidence", "nephrology has documented" in p3_golden)
check("Golden P3: names the alternative (acetaminophen)", "acetaminophen" in p3_golden)
lecture_words = ["it is essential", "it is critical", "it is imperative", "must be emphasized", "cannot be overstated"]
p3_lecture = [w for w in lecture_words if w in p3_golden.lower()]
check(f"Golden P3: no lecturing tone (found: {p3_lecture})", len(p3_lecture) == 0)
# Check it doesn't perform the task's synthesis (clinical-voice CAUTION)
synthesis_words = ["central", "primary concern", "most important", "key finding", "critical issue"]
p3_synth = [w for w in synthesis_words if w in p3_golden.lower()]
check(f"Golden P3: no synthesis/weighting language (found: {p3_synth})", len(p3_synth) == 0)

# =====================================================================
print("\n" + "=" * 70)
if fails:
    print(f"RESULT: {len(fails)} FAILURES, {len(warns)} WARNINGS")
    for f in fails:
        print(f"  FAIL: {f}")
else:
    print(f"RESULT: ALL CHECKS PASSED ({len(warns)} warnings)")
for w in warns:
    print(f"  WARN: {w}")
print("=" * 70)
