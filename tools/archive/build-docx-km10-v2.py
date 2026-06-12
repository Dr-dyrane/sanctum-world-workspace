#!/usr/bin/env python3
"""KM10 v2 golden - Mode A text swap of the v1 golden (reseed per Abi 6/11).

Only query item 2 (metabolic encephalopathy) changes: the decline is rewritten in the
attending's own voice, anchored on his team's documented assessment, and rests on clinical
grounds rather than on the diagnosis being added after discharge. Items 1 and 3, all chrome,
and styles.xml are left byte-identical (Mode A). Locators are content-based, not positional.
"""
import os, sys, zipfile, shutil
from docx import Document

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SRC = os.path.join(ROOT, "worlds/korvin-merrow/task-setup/platform/task10/archive/2026-06-11-v1-reseed-abi/golden-KM10-v1.docx")
OUT = os.path.join(ROOT, "worlds/korvin-merrow/task-setup/platform/task10/current/golden-KM10-v2.docx")

ITEM2_HEADER = "Item 2 - metabolic encephalopathy: declined on clinical grounds"
ITEM2_BODY = (
    "My team assessed and documented the patient's altered mental status throughout the admission. "
    "We characterized it as altered baseline mental status with functional decline, explicitly multifactorial, "
    "with infection physiology, dehydration, medication effects, and limited baseline reserve all contributing, "
    "and we did not identify a single metabolic derangement as its cause. It was mild and intermittent, "
    "predominantly evening confusion, and it improved steadily as the infection, volume status, and renal function "
    "recovered. By hospital day 3 the cognitive findings were reframed and managed as a functional and cognitive "
    "readiness issue through physical and occupational therapy, with executive slowing under complexity rather than "
    "an acute neurologic syndrome. No encephalopathy-directed workup was undertaken, no neurology consultation was "
    "obtained, and no treatment was directed at an encephalopathy. The indicators the query cites are real, but in "
    "this course they reflect a multifactorial, improving, symptom-level alteration in a medically complex patient, "
    "not a distinct toxic-metabolic encephalopathy. I am keeping the documentation at the level my team assessed and "
    "recorded. If the program wishes to capture cognitive and functional risk, the supported vehicle is the documented "
    "functional and cognitive findings, not an encephalopathy diagnosis the contemporaneous assessment does not establish."
)

shutil.copyfile(SRC, OUT)
d = Document(OUT)

swapped_header = swapped_body = False
for p in d.paragraphs:
    t = p.text.strip()
    if t.startswith("Item 2 - metabolic encephalopathy"):
        p.runs[0].text = ITEM2_HEADER          # single run; preserves formatting (Mode A)
        for r in p.runs[1:]:
            r.text = ""
        swapped_header = True
    elif t.startswith("I am not adding this diagnosis. The contemporaneous record documents intermittent confusion"):
        p.runs[0].text = ITEM2_BODY
        for r in p.runs[1:]:
            r.text = ""
        swapped_body = True

assert swapped_header and swapped_body, f"locators failed: header={swapped_header} body={swapped_body}"

# metadata scrub
cp = d.core_properties
cp.author = ""
cp.last_modified_by = ""
cp.title = ""
cp.comments = ""
cp.category = ""
d.save(OUT)

# integrity gate
raw = open(OUT, "rb").read()
assert raw.find(b"PK\x05\x06") >= 0, "no EOCD - truncated zip"
names = zipfile.ZipFile(OUT).namelist()
assert "word/styles.xml" in names and "word/document.xml" in names, "missing parts"
Document(OUT)  # must reopen

# styles.xml byte-identical vs source (Mode A invariant)
s_src = zipfile.ZipFile(SRC).read("word/styles.xml")
s_out = zipfile.ZipFile(OUT).read("word/styles.xml")
assert s_src == s_out, "styles.xml drifted - not a clean Mode A edit"

# banned characters
body_txt = "\n".join(p.text for p in d.paragraphs)
for bad in ("—", "–", "→"):
    assert bad not in body_txt, f"banned char {bad!r} present"

print("gate ok | header+body swapped | styles.xml byte-identical | no em/en/arrow")
print("OUT:", OUT)
