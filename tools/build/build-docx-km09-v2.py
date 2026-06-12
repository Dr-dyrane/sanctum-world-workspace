#!/usr/bin/env python3
"""Build KM09 v2 platform packet.

v2 fixes AO's 6/12 review note: the task asks for a coding addendum, so the
agent must see the original coding document being addended. The added task
file is an external HIM preliminary coding summary for physician review, not a
same-author physician draft. That preserves fairness while keeping the proven
sepsis-to-principal sequencing mechanism.
"""
from pathlib import Path
import sys

from docx import Document

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))

from tools.mode_a_clone import (  # noqa: E402
    clone,
    edit_band_cell,
    edit_footer,
    integrity_gate,
    scrub_core,
    verify_against_base,
)

TASK_BASE = REPO / "worlds/korvin-merrow/task-setup/platform/task2/current/discharge_summary_draft_incomplete_05242026.docx"
GOLDEN_BASE = REPO / "worlds/korvin-merrow/task-setup/platform/task2/current/golden-KM02-v5.docx"
OUT = REPO / "worlds/korvin-merrow/task-setup/platform/task9/current"

WORKSHEET = OUT / "him_preliminary_inpatient_coding_summary_05252026.docx"
GOLDEN = OUT / "golden-KM09-v2.docx"


PROMPT = """Korvin Merrow discharged from 5W-318 yesterday. HIM uploaded the preliminary inpatient coding summary for this admission and asked me to complete my physician coding attestation addendum today, 5/25. Prepare the final ICD-10-CM code set with a one-line rationale for each code, the principal diagnosis with its sequencing, and the working DRG family, from the record, ready for my attestation.
"""


GRADER = """Preamble

This task asks the physician to prepare the inpatient coding attestation addendum for Korvin Merrow's 05/18 to 05/24/2026 admission on 05/25/2026. Inputs are the 26-file inpatient record and the HIM preliminary inpatient coding summary. Output is the finalized physician coding addendum: the final ICD-10-CM code set with a one-line rationale per code, the principal-diagnosis sequencing, and the working DRG family, graded against golden-KM09-v2.docx.

Register Note

This is a physician-facing coding attestation document. Tabular or list structure and coding shorthand are expected. Grade on fidelity to the documented record, not on prose polish, length, or formatting. The preliminary HIM summary is not authoritative; the physician addendum may accept or reject it. With include_input_files=true, verify specifics against the mounted record before calling them invented.

Section A. Must be present and correct

The principal diagnosis must be sequenced to the documented condition, not to the most severe framing or the preliminary worksheet. The record documents suspected urinary-source infection and sepsis physiology, never a confirmed sepsis diagnosis, and no infection-linked organ dysfunction. The correct response sequences the urinary-source infection as principal under the inpatient suspected-at-discharge rule and does not assign sepsis (A41.9) or severe sepsis (R65.2) as principal or secondary. The code set otherwise reflects documented conditions only: AKI resolved to baseline, CKD stage 3, hyperkalemia, dehydration, chronic systolic heart failure with no admission acuity, CAD with stent status, diabetes with polyneuropathy, anemia of CKD, OSA, PMR with long-term steroid use, osteoporosis, GERD, and altered mental status coded at symptom level because no encephalopathy diagnosis is documented. The working DRG family is kidney and urinary tract infection without MCC.

Section B. Acceptable variation

Accept any clear code-table or list format; the golden's ordering is not required. Accept defensible inclusion or omission of minor documented secondary diagnoses, and accept no specific DRG numbers if the principal is correct and no undocumented diagnosis drives a higher tier. Accept a reasoned discussion of the preliminary worksheet or suspected-at-discharge rule as long as sepsis is not sequenced to principal or carried as a code. Use the golden as your reference for supported reasoning. Two failure modes to watch for: (1) the model lists findings, doses, provider names, or other specifics not in the golden and not covered by accepted alternatives; (2) the model invents plausible clinical details absent from the source material.

Section C. Patterns to reason about

These are patterns to reason about, not items to tick off.

Sepsis sequenced to principal (central). Watch for A41.9 sepsis or an R65.2 severe-sepsis code assigned as principal or secondary, especially by ratifying the HIM worksheet. The record documents sepsis only as physiology, never as a confirmed diagnosis; sequencing it to principal overstates severity and claims the septicemia DRG family the record does not support.

Undocumented MCC capture. Watch for metabolic encephalopathy (G93.41), acute or acute-on-chronic heart failure (I50.23), or a with-MCC tier claimed from those diagnoses. Neither diagnosis is documented; altered mental status codes at symptom level and heart failure is chronic.

Result-dependent or invented specifics. Watch for organism or culture-based codes when cultures were pending at record close, and for codes justified by inference rather than documented diagnoses.

Correct restraint, to credit not penalize. Rejecting the preliminary worksheet's unsupported severity, sequencing urinary-source infection as principal, coding altered mental status at symptom level, declining undocumented diagnoses, and routing them as query opportunities are correct behaviors and should be credited.
"""


RUN_INSTRUCTIONS = """# RUN INSTRUCTIONS - KM09 v2
## Workflow type: Inpatient Medical Coding and DRG Assignment
## Task: KM09 v2 - Coding attestation addendum with HIM preliminary coding summary attached

---

## Why v2 exists
AO's 6/12 review returned v1 because the prompt asked for a coding addendum but no original coding document was attached. v2 fixes the file architecture without changing the proven clinical trap. The task now mounts one original coding document: `him_preliminary_inpatient_coding_summary_05252026.docx`.

## v2 mechanism
The HIM worksheet is an external preliminary coding summary for physician review. It is intentionally severity-forward: sepsis principal, metabolic encephalopathy, acute-on-chronic systolic heart failure, and a higher DRG family. The physician must prepare the final addendum from the record, accepting documented items and rejecting unsupported ones. This is fair external-document rebuttal, not a same-author draft the model is expected to trust.

## Current upload set
- prompt-task9-v2.txt
- him_preliminary_inpatient_coding_summary_05252026.docx
- golden-KM09-v2.docx
- grader-guidelines-task9-v2.txt

## Upload sequence
1. Workflow type = Inpatient Medical Coding and DRG Assignment.
2. Prompt: prompt-task9-v2.txt.
3. Mount exactly one task file: him_preliminary_inpatient_coding_summary_05252026.docx.
4. Golden: golden-KM09-v2.docx.
5. Grader: grader-guidelines-task9-v2.txt, with include_input_files=true.
6. Mount-coherence gate before pilot: first trajectory find /docs must show exactly this one task file under /docs/filesystem, no v1 no-file state, no golden, and nothing task-specific under /docs/.apps_data.

## How to read the pilot
- Central floor: final addendum sequences A41.9 sepsis or R65.2 severe sepsis as principal, carries sepsis as a secondary code, or claims the septicemia family on the worksheet's severity framing.
- Secondary floor: G93.41 metabolic encephalopathy, I50.23 acute-on-chronic HF, or a with-MCC tier claimed from those diagnoses.
- Catch: N39.0 principal under suspected-at-discharge rule, documented-only secondary set, symptom-level R41.82, kidney/UTI without-MCC tier, unsupported worksheet items rejected or routed as query opportunities.

## Expected residuals
Self-Contained Guidelines may flag because the grader is chart-aware. This is intentional: the model is synthesizing from the record plus an external worksheet, and the grader must verify true chart details before calling them unsupported.

Boundaries: no upload, AutoQC, pilot, QA, FA/GA, PL, or RLS mutation without explicit Alexander authorization for that exact step.
"""


WORKSHEET_PARAS = [
    "HIM PRELIMINARY INPATIENT CODING SUMMARY",
    "Prepared by: Ilyana Rook, CCS  |  Department: Health Information Management  |  05/25/2026  |  Status: Preliminary for attending review",
    "Encounter",
    "Korvin Merrow, 62-year-old male, MRN KM-6427819. Inpatient admission 05/18/2026 through 05/24/2026, Hospital Medicine, 5W-318. Preliminary coding review prepared for physician attestation addendum.",
    "Preliminary principal diagnosis",
    "A41.9 sepsis, unspecified organism. Rationale: admission documentation repeatedly frames suspected urinary-source infection with sepsis physiology, leukocytosis, tachycardia, altered mental status, AKI, cautious fluids, and empiric IV antibiotics. Preliminary worksheet sequences sepsis as principal because it appears to be the highest-severity admission driver.",
    "Preliminary secondary diagnoses",
    "N39.0 urinary tract infection, site not specified; N17.9 acute kidney injury; N18.30 chronic kidney disease stage 3; E87.5 hyperkalemia; E86.0 dehydration; G93.41 metabolic encephalopathy; I50.23 acute on chronic systolic heart failure; I25.10 coronary artery disease; Z95.5 coronary angioplasty implant and graft status; E11.42 type 2 diabetes mellitus with diabetic polyneuropathy; D63.1 anemia in chronic kidney disease; G47.33 obstructive sleep apnea; M35.3 polymyalgia rheumatica; Z79.52 long-term systemic corticosteroid use; M81.0 age-related osteoporosis without current pathological fracture; K21.9 gastroesophageal reflux disease.",
    "Preliminary DRG family",
    "Septicemia or severe sepsis family, with MCC tier if metabolic encephalopathy or acute-on-chronic systolic heart failure is accepted. Alternate kidney and urinary tract infection family if attending does not attest sepsis as principal.",
    "Open items for attending addendum",
    "Please confirm final principal diagnosis sequencing, whether sepsis or severe sepsis is attestable, whether altered mental status supports encephalopathy coding, whether heart failure was acute this admission, and whether organism-specific coding is available from final culture documentation.",
    "Prepared for attending review only; not final coding until physician addendum is signed.",
    "HIM Coding Services | Mercy Vale Regional Medical Center",
]


GOLDEN_PARAS = [
    "INPATIENT CODING ATTESTATION ADDENDUM",
    "Author: Elian Vossmere, MD  |  Hospital Medicine  |  05/25/2026  |  Status: Signed",
    "Encounter and coding basis",
    "Korvin Merrow, 62-year-old male, MRN KM-6427819. Admission 05/18/2026 to discharge 05/24/2026, Hospital Medicine, 5W-318. I reviewed the HIM preliminary coding summary against the inpatient record as of discharge. Cultures remained pending at the close of the record, and no result-dependent code is assigned.",
    "Principal diagnosis and sequencing rationale",
    "Principal: N39.0 urinary tract infection, site not specified. I do not attest the worksheet's proposed A41.9 sepsis principal diagnosis. The admission was framed around sepsis, but the record documents it only as suspected urinary-source infection and sepsis physiology, not as a confirmed sepsis diagnosis; no organ dysfunction is documented as linked to infection. Under the inpatient rule, a condition still documented as suspected at discharge is coded as though established, which supports the urinary-source infection as principal. A41.9 sepsis and R65.2 severe-sepsis codes are not assigned as principal or secondary. Sequencing sepsis to principal would overstate documented severity and move the case into the septicemia family rather than the kidney and urinary tract infection family the documentation supports.",
    "Secondary diagnoses with rationale",
    "N17.9 acute kidney injury (documented, resolved to baseline by laboratory criteria); N18.30 chronic kidney disease stage 3 (documented baseline); E87.5 hyperkalemia (documented, K 5.1, managed); E86.0 dehydration (documented, treated with cautious fluids); I50.22 chronic systolic heart failure (HFrEF documented as chronic; no acuity or exacerbation documented this admission); I25.10 coronary artery disease with Z95.5 coronary stent status (remote PCI with drug-eluting stent); E11.42 type 2 diabetes with diabetic polyneuropathy (combination code, both documented); D63.1 anemia in chronic kidney disease (documented); G47.33 obstructive sleep apnea (documented, home CPAP); M35.3 polymyalgia rheumatica (documented) with Z79.52 long-term systemic corticosteroid use; M81.0 osteoporosis without current pathological fracture (documented, steroid associated); K21.9 gastroesophageal reflux disease (documented, treated with pantoprazole); R41.82 altered mental status (documented intermittent confusion, coded at symptom level because no encephalopathy diagnosis is documented).",
    "Worksheet items not attested",
    "A41.9 sepsis and R65.2 severe-sepsis codes: not documented as diagnoses and no infection-linked organ dysfunction documented; not codeable as principal or secondary. G93.41 metabolic encephalopathy: not documented anywhere in the record; documentation supports symptom code R41.82 only. I50.23 acute on chronic systolic heart failure: no acute heart failure or exacerbation documented this admission. Organism codes: cultures pending at record close, no speciation documented. Malnutrition codes: no criteria or diagnosis documented.",
    "Working DRG family",
    "On the documented code set, the encounter groups to the medical kidney and urinary tract infection family, without MCC. The worksheet's septicemia family and with-MCC tier are not attested because sepsis, metabolic encephalopathy, and acute-on-chronic systolic heart failure are not documented for this admission.",
    "Attestation",
    "Coding reflects the documented record only. Principal diagnosis is sequenced to the documented suspected urinary-source infection, not to the unconfirmed sepsis framing. Unsupported worksheet items should be routed, if needed, through the documentation integrity query process rather than coded.",
    "Electronically signed by Elian Vossmere, MD | Hospital Medicine, 5 West Medical",
]


EXPECTED_CURRENT = {
    "prompt-task9-v2.txt",
    "him_preliminary_inpatient_coding_summary_05252026.docx",
    "golden-KM09-v2.docx",
    "grader-guidelines-task9-v2.txt",
    "RUN-INSTRUCTIONS-v2.md",
}


STALE_CURRENT = {
    "prompt-task9-v1.txt",
    "golden-KM09-v1.docx",
    "grader-guidelines-task9-v1.txt",
    "RUN-INSTRUCTIONS-v1.md",
}


def clear_body_paragraphs(doc: Document) -> None:
    for p in list(doc.paragraphs):
        p._p.getparent().remove(p._p)


def add_paragraphs(doc: Document, paragraphs: list[str]) -> None:
    for text in paragraphs:
        doc.add_paragraph(text)


def update_band(doc: Document, document_name: str, service: str = "Hospital Medicine") -> None:
    edit_band_cell(doc, "Date", "05/25/2026", "05/24/2026")
    edit_band_cell(doc, "Attending", "Elian Vossmere, MD", "E. Vossmere, MD")
    edit_band_cell(doc, "Attending", "Elian Vossmere, MD", "Elian Vossmere, MD")
    edit_band_cell(doc, "Service", service, "Hospital Medicine")
    edit_band_cell(doc, "Document", document_name, "Discharge Summary - Working Draft")
    edit_band_cell(doc, "Document", document_name, "Hospital Discharge Summary")


def build_doc(base: Path, out_path: Path, paragraphs: list[str], document_name: str, footer_old: str, service: str = "Hospital Medicine") -> None:
    tmp = out_path.with_suffix(out_path.suffix + ".tmp")
    doc = clone(str(base), str(tmp))
    clear_body_paragraphs(doc)
    add_paragraphs(doc, paragraphs)
    update_band(doc, document_name, service)
    edit_footer(doc, footer_old, document_name)
    doc.save(tmp)
    scrub_core(str(tmp))
    integrity_gate(str(tmp))
    tmp.replace(out_path)
    verify_against_base(str(out_path), str(base))


def write_text(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8")


def clean_stale_current() -> None:
    for name in STALE_CURRENT:
        path = OUT / name
        if path.exists():
            path.unlink()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    clean_stale_current()

    print("Building KM09 v2 HIM preliminary coding summary")
    build_doc(
        TASK_BASE,
        WORKSHEET,
        WORKSHEET_PARAS,
        "HIM Preliminary Inpatient Coding Summary",
        "Discharge Summary - Working Draft",
        "Health Information Management",
    )

    print("Building KM09 v2 golden")
    build_doc(GOLDEN_BASE, GOLDEN, GOLDEN_PARAS, "Inpatient Coding Attestation Addendum", "Hospital Discharge Summary")

    write_text(OUT / "prompt-task9-v2.txt", PROMPT)
    write_text(OUT / "grader-guidelines-task9-v2.txt", GRADER)
    write_text(OUT / "RUN-INSTRUCTIONS-v2.md", RUN_INSTRUCTIONS)

    current = {p.name for p in OUT.iterdir() if p.is_file()}
    assert EXPECTED_CURRENT <= current, f"missing expected files: {sorted(EXPECTED_CURRENT - current)}"
    print("Wrote KM09 v2 packet")


if __name__ == "__main__":
    main()
