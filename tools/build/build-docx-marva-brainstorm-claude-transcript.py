#!/usr/bin/env python3
"""Build the Marva Lydell Brainstorm transcript DOCX."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))

from tools.mode_a_clone import integrity_gate, scrub_core


SOURCE = REPO / "worlds/marva-lydell/submission/Marva_Lydell_Brainstorm_Claude_Transcript.md"
OUT = REPO / "worlds/marva-lydell/submission/Marva_Lydell_Brainstorm_Claude_Transcript.docx"
BANNED_RE = re.compile("[\u2013\u2014\u2190-\u21ff]")
CONTENT_WIDTH_DXA = 9360

ROLE_COLORS = {
    "System": "6B7280",
    "Human": "1F4D78",
    "Assistant": "8A5A44",
}

ROLE_FILLS = {
    "System": "F8FAFC",
    "Human": "F6F8FA",
    "Assistant": "FBF7F0",
}


def set_paragraph_spacing(paragraph, before=0, after=6, line=280) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:before"), str(before * 20))
    spacing.set(qn("w:after"), str(after * 20))
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "auto")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=150, bottom=110, end=150) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color: str, accent: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "right", "bottom"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "14")
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), accent)


def set_table_width(table, width: int) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width))
    tbl_w.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    col = OxmlElement("w:gridCol")
    col.set(qn("w:w"), str(width))
    grid.append(col)
    for row in table.rows:
        cell = row.cells[0]
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_w = tc_pr.find(qn("w:tcW"))
        if tc_w is None:
            tc_w = OxmlElement("w:tcW")
            tc_pr.append(tc_w)
        tc_w.set(qn("w:w"), str(width))
        tc_w.set(qn("w:type"), "dxa")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_plain(doc: Document, text: str, bold_prefix: bool = False):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=5, line=270)
    if bold_prefix and ":" in text:
        prefix, rest = text.split(":", 1)
        run = p.add_run(prefix + ":")
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run = p.add_run(rest)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
    else:
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
    return p


def add_turn_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=14, after=1, line=250)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("1F4D78")


def add_role_label(doc: Document, role: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=3, after=3, line=250)
    run = p.add_run(role)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(ROLE_COLORS.get(role, "1F4D78"))


def add_code_block(doc: Document, block_lines: list[str], role: str | None) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, CONTENT_WIDTH_DXA)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, ROLE_FILLS.get(role or "", "F6F8FA"))
    set_cell_borders(cell, "D0D7DE", ROLE_COLORS.get(role or "", "6B7280"))
    set_cell_margins(cell)
    for index, block_line in enumerate(block_lines or [""]):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        set_paragraph_spacing(p, after=1, line=232)
        run = p.add_run(block_line if block_line else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(8.7)
        run.font.color.rgb = RGBColor.from_string("111827")
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, after=3, line=200)


def build_docx() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    if BANNED_RE.search(text):
        raise ValueError("Source contains banned dash or arrow characters")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = RGBColor.from_string("2E74B5")

    current_role: str | None = None
    in_code = False
    code_lines: list[str] = []

    for raw in text.splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines, current_role)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
            continue
        if re.fullmatch(r"\d+\.", line.strip()):
            add_turn_number(doc, line.strip())
            current_role = None
            continue
        if line.strip() in {"System", "Human", "Assistant"}:
            current_role = line.strip()
            add_role_label(doc, current_role)
            continue
        add_plain(doc, line, bold_prefix=":" in line and current_role is None)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Marva Lydell Brainstorm transcript"
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("6B7280")

    doc.save(str(OUT))
    scrub_core(str(OUT))
    integrity_gate(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build_docx()
