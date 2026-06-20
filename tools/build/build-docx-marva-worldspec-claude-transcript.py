#!/usr/bin/env python3
"""Build the Marva Lydell World Spec Claude transcript DOCX."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
import sys as _sys; _sys.path.insert(0, str(REPO))
from tools.mode_a_clone import make_deterministic
SOURCE = REPO / "worlds/marva-lydell/submission/Marva_Lydell_World_Spec_Claude_Transcript.md"
OUT = REPO / "worlds/marva-lydell/submission/Marva_Lydell_World_Spec_Claude_Transcript.docx"

BANNED_RE = re.compile("[\u2013\u2014\u2190-\u21ff]")
CONTENT_WIDTH_DXA = 9360

ROLE_COLORS = {
    "System": "6B7280",
    "Human": "1F4D78",
    "Assistant": "8A5A44",
}

ROLE_FILLS = {
    "System": "F8FAFC",
    "Human": "F6F8FA",
    "Assistant": "FBF7F0",
}

ROLE_BORDERS = {
    "System": "CBD5E1",
    "Human": "D0D7DE",
    "Assistant": "E4D8C8",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color: str, accent: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "right", "bottom"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "14")
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), accent or color)


def set_table_widths(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph_spacing(paragraph, before=0, after=6, line=280) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:before"), str(before * 20))
    spacing.set(qn("w:after"), str(after * 20))
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "auto")


def add_para(doc, text: str, style: str | None = None, bold_prefix: bool = False):
    p = doc.add_paragraph(style=style)
    set_paragraph_spacing(p, after=6, line=280)
    if bold_prefix and ":" in text:
        prefix, rest = text.split(":", 1)
        run = p.add_run(prefix + ":")
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(11)
        if rest:
            run = p.add_run(rest)
            run.font.name = "Arial"
            run.font.size = Pt(11)
    else:
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(11)
    return p


def add_meta_line(doc, text: str):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=3, line=250)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string("374151")
    return p


def add_turn_number(doc, text: str):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=14, after=1, line=250)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string("1F4D78")
    run.bold = True


def add_role_label(doc, role: str):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=3, after=3, line=250)
    run = p.add_run(role)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(ROLE_COLORS.get(role, "1F4D78"))
    run.bold = True


def add_code_block(doc, lines: list[str], role: str | None) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_widths(table, [CONTENT_WIDTH_DXA])
    cell = table.rows[0].cells[0]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    fill = ROLE_FILLS.get(role or "", "F6F8FA")
    border = ROLE_BORDERS.get(role or "", "D0D7DE")
    accent = ROLE_COLORS.get(role or "", "6B7280")
    set_cell_shading(cell, fill)
    set_cell_borders(cell, border, accent)
    set_cell_margins(cell, top=130, start=170, bottom=130, end=170)

    block_lines = lines or [""]
    for idx, block_line in enumerate(block_lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        set_paragraph_spacing(p, after=1, line=232)
        run = p.add_run(block_line if block_line else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(8.7)
        run.font.color.rgb = RGBColor.from_string("111827")

    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, after=3, line=200)
    doc.add_paragraph()


def build_docx() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    if BANNED_RE.search(text):
        raise ValueError("Source contains banned dash or arrow characters")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    for style_name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    lines = text.splitlines()
    title = lines[0].removeprefix("# ").strip()
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(title_p, after=5, line=280)
    run = title_p.add_run(title)
    run.font.name = "Arial"
    run.font.size = Pt(19)
    run.font.color.rgb = RGBColor.from_string("111827")
    run.bold = True

    skip_first = True
    in_code = False
    code_lines: list[str] = []
    current_role: str | None = None
    seen_turn = False
    for raw in lines:
        line = raw.strip()
        if skip_first:
            skip_first = False
            continue
        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines, current_role)
                code_lines = []
                in_code = False
            else:
                in_code = True
                code_lines = []
            continue
        if not line:
            if in_code:
                code_lines.append("")
            continue
        if in_code:
            code_lines.append(raw.rstrip())
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(line[3:], style="Heading 1")
            set_paragraph_spacing(p, before=16, after=8, line=280)
        elif line.startswith("### "):
            p = doc.add_paragraph(line[4:], style="Heading 2")
            set_paragraph_spacing(p, before=12, after=6, line=280)
        elif re.match(r"^\d+\.$", line):
            add_turn_number(doc, line)
            seen_turn = True
        elif line in {"System", "Human", "Assistant"}:
            current_role = line
            add_role_label(doc, line)
        elif re.match(r"^\d+\. ", line):
            add_para(doc, line)
        elif line.startswith("Alexander:") or line.startswith("Claude:"):
            add_para(doc, line, bold_prefix=True)
        elif not seen_turn and ":" in line:
            add_meta_line(doc, line)
        else:
            add_para(doc, line)

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_run = footer.add_run("Marva Lydell World Spec Claude transcript")
        footer_run.font.name = "Arial"
        footer_run.font.size = Pt(8)
        footer_run.font.color.rgb = RGBColor.from_string("666666")

    doc.core_properties.author = ""
    doc.core_properties.comments = ""
    doc.core_properties.subject = ""
    doc.core_properties.title = "Marva Lydell World Spec Claude Transcript"
    doc.core_properties.keywords = ""
    doc.core_properties.last_modified_by = ""
    doc.save(OUT)

    raw = OUT.read_bytes()
    assert raw.find(b"PK\x05\x06") >= 0, "DOCX missing EOCD"
    with zipfile.ZipFile(OUT) as zf:
        names = set(zf.namelist())
        assert "word/document.xml" in names and "word/styles.xml" in names
        xml = zf.read("word/document.xml").decode("utf-8")
        if BANNED_RE.search(xml):
            raise ValueError("DOCX XML contains banned dash or arrow characters")
    Document(str(OUT))
    make_deterministic(str(OUT))


if __name__ == "__main__":
    build_docx()
