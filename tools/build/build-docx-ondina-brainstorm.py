#!/usr/bin/env python3
"""Build the Ondina Vasquell Brainstorm DOCX via Mode A clone."""

from __future__ import annotations

import copy
import tempfile
import sys
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))

from tools.mode_a_clone import integrity_gate, scrub_core, set_text, verify_against_base


BASE = REPO / "worlds/korvin-merrow/submission/Korvin_Merrow_Brainstorm.docx"
OUT = REPO / "worlds/ondina-vasquell/submission/Ondina_Vasquell_Brainstorm.docx"


def remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._p.getparent().remove(paragraph._p)


def fill_cell(cell, lines: list[str]) -> None:
    """Replace a table cell with paragraph clones from its first paragraph."""
    if not lines:
        lines = [""]
    proto = copy.deepcopy(cell.paragraphs[0]._p)
    for paragraph in list(cell.paragraphs):
        remove_paragraph(paragraph)
    for text in lines:
        new_el = copy.deepcopy(proto)
        new_p = Paragraph(new_el, cell)
        set_text(new_p, text)
        cell._tc.append(new_el)


def lines(*items: str | None) -> list[str]:
    return ["" if item is None else item for item in items]


WORLD_SETUP = lines(
    "World Type: Typical Clinical World",
    None,
    "This is an inpatient hospital medicine world: a 6-day admission for a limb-threatening diabetic foot infection in Ondina Vasquell, a 68-year-old Spanish-preferred woman with long-standing insulin-dependent type 2 diabetes complicated by diabetic peripheral neuropathy and mild diabetic retinopathy, CKD stage 3b with anemia of CKD, peripheral arterial disease, HFpEF, hypertension, dyslipidemia, obesity, obstructive sleep apnea, knee osteoarthritis, and limited mobility. She presents after roughly three weeks of outpatient wound deterioration. The course runs ED presentation, admission for IV antibiotics, podiatry soft-tissue debridement, vascular evaluation with perfusion questions still open, and steady improvement in infection markers. She lives in a second-floor walk-up; her adult daughter helps but works nights; insurance is Medicare Advantage with Medicaid secondary.",
    None,
    "The world snapshot closes on May 21, 2026 at 18:00, at the point of maximal realistic tension: medically improving but operationally unsafe. Perfusion is not fully resolved, offloading is not reliably teach-backed, the home setup is poor, and payer, SNF, and DME issues are all active. Every task is an independent encounter anchored strictly after the snapshot: discharge-day work on May 22, administrative and payer surfaces on May 24 to May 26, and quality, referral, and safety-review work on June 4 to June 11. Each task can be answered from the planned chart, payer, pharmacy, and quality materials available to that requester.",
)

FRICTIONS = lines(
    "1. Hospital medicine versus the Medicare Advantage medical director",
    "The payer reads improving markers after debridement as ready for home with home health; the treating team reads the same chart as unsafe without SNF-level wound care, offloading training, and equipment in place.",
    None,
    "2. Daughter versus PT/OT and case management",
    "The daughter wants her mother home and offers real but night-shift-limited support; PT/OT findings, the second-floor walk-up, and wound-care frequency say home is not yet safe.",
    None,
    "3. Podiatry versus vascular surgery",
    "Debride, offload, and advance wound care now versus settle perfusion adequacy and revascularization sequencing before setting healing expectations.",
    None,
    "4. HIM/CDI versus the treating team",
    "Severity-capture pressure toward acute osteomyelitis and POA specificity versus the treating clinicians' documented restraint.",
    None,
    "5. Pharmacy benefit manager versus infectious disease",
    "Formulary substitution pressure versus a renal-safe antibiotic regimen at her eGFR.",
)

TRAPS = lines(
    "1. Renal antibiotic dosing under a shifting eGFR",
    "Type: contradictory and buried information. The MAR, creatinine trend, and ID note contradict a dose carried from admission renal function, and the tempting formulary substitute is unsafe at her eGFR. World-level substrate; task-level surfaces force it in the med rec and pharmacy rejection tasks.",
    None,
    "2. Equivocal osteomyelitis with unsupported specificity",
    "Type: insufficient and uncertain information. MRI shows marrow edema and says early osteomyelitis cannot be excluded; podiatry documents no exposed bone; no bone specimen confirms osteomyelitis; ID treats deep diabetic foot infection without signing acute osteomyelitis. Supported coding is diabetic foot ulcer with cellulitis or deep soft tissue infection; acute osteomyelitis requires treating-clinician clarification. World-level substrate; task-level HIM and CDI documents pressure the clinician toward over-specificity.",
    None,
    "3. Diabetic foot ulcer versus pressure injury code family, plus POA status",
    "Type: source-of-truth ambiguity across wound care nursing, podiatry, and nursing skin assessments. World-level substrate; the coding worksheet and abstraction fields force the distinction.",
    None,
    "4. Vascular adequacy overstated",
    "Type: contradictory EMR information. A palpable-pulse or Doppler-signal note reads reassuring while ankle-brachial and toe-brachial values with noncompressible vessels and the vascular consult keep perfusion genuinely open. World-level substrate; payer, determination, and referral tasks force a stance.",
    None,
    "5. Offloading, teach-back, stairs, and DME readiness gaps",
    "Type: buried significant information. The signal is spread across PT/OT evaluations, nursing teaching notes, and case management; no single source states the disposition conclusion. World-level substrate; the payer appeal, determination, safety review, and discharge-instruction tasks each force a different deliverable-level decision.",
    None,
    "6. Quiet quality-measure lookback disqualifier",
    "Type: temporal complexity. One date or exclusion in the outpatient record contradicts naive numerator or denominator capture. Task-level abstraction fields force the value, exclusion, or unable-to-determine stance.",
    None,
    "7. Culture provenance hierarchy",
    "Type: source-of-truth ambiguity. Superficial swab, deep tissue, and blood culture results carry different authority for antibiotic rationale. World-level substrate; med rec, CDI, and pharmacy tasks test whether the clinician weights culture sources correctly.",
    None,
    "8. SDOH context",
    "Language preference, night-shift caregiver, and walk-up housing are deliberately load-bearing context across discharge-safety tasks but never themselves the central trap.",
)

TASKS = lines(
    "Ten tasks, each an independent post-snapshot encounter, each mapped to an approved tracker workflow and named with its requester, structure, forced slot, task-level trap, and anchor. Ten distinct workflows are intentional: consolidating them would blur real, distinct physician-facing deliverables, and each task still names its structural category and forcing function so variety is visible at a glance.",
    None,
    "1. Discharge medication reconciliation safety table",
    "Workflow mapping: Discharge Medication Reconciliation.",
    "Requester and anchor: Hospitalist attending, May 22, 2026 at 08:30.",
    "Structure and forced slot: Forced inventory; continue, hold, change, stop, or defer per row.",
    "Task-level trap: Preliminary discharge med list carries admission renal-dose logic despite later eGFR and ID/pharmacy notes.",
    None,
    "2. Physician coding attestation against the HIM preliminary worksheet",
    "Workflow mapping: Inpatient Medical Coding and DRG Assignment.",
    "Requester and anchor: HIM coding lead, May 22, 2026 at 09:00.",
    "Structure and forced slot: Forced inventory; principal diagnosis, POA, and code family per row.",
    "Task-level trap: HIM worksheet pressures pressure-injury family and acute osteomyelitis POA without treating/pathologic support.",
    None,
    "3. Attending response to a CDI query",
    "Workflow mapping: Clinical Documentation Improvement (CDI) Query Response Review.",
    "Requester and anchor: CDI specialist, May 24, 2026 at 10:00.",
    "Structure and forced slot: External ratify-or-refute; agree, decline, or unable to determine per item.",
    "Task-level trap: CDI query asks for acute osteomyelitis specificity and severity language that the treating record does not establish.",
    None,
    "4. Physician appeal of the Medicare Advantage SNF-authorization denial",
    "Workflow mapping: Claims Denial Analysis and Appeal Preparation.",
    "Requester and anchor: Hospitalist attending, at case management request, May 24, 2026 at 15:00.",
    "Structure and forced slot: External ratify-or-refute; appeal, accept, or narrow.",
    "Task-level trap: Denial frames improving infection markers as home-with-home-health readiness while omitting offloading, stairs, caregiver, and perfusion barriers.",
    None,
    "5. Response to a pharmacy insurance claim rejection",
    "Workflow mapping: Pharmacy Insurance Claim Rejection Resolution.",
    "Requester and anchor: Inpatient pharmacist, May 25, 2026 at 09:00.",
    "Structure and forced slot: External ratify-or-refute; substitute, appeal, hold, or exception request.",
    "Task-level trap: PBM-preferred substitute is administratively easy but unsafe against renal function, culture hierarchy, or interaction context.",
    None,
    "6. Continued-stay determination",
    "Workflow mapping: Utilization Review Concurrent Stay Documentation.",
    "Requester and anchor: Physician advisor, May 26, 2026 at 11:00.",
    "Structure and forced slot: Determination; binding continued-stay verdict.",
    "Task-level trap: Review note treats post-debridement improvement as level-of-care readiness despite unresolved operational limb-safety barriers.",
    None,
    "7. Diabetes quality-measure chart abstraction",
    "Workflow mapping: HEDIS Medical Record Chart Abstraction and Review.",
    "Requester and anchor: Quality abstraction nurse, June 4, 2026 at 09:00.",
    "Structure and forced slot: Extraction to schema; value, exclusion, or unable to determine per field.",
    "Task-level trap: A quiet lookback date or exclusion in the outpatient record contradicts naive diabetes-measure capture.",
    None,
    "8. Vascular surgery referral letter with required disposition table",
    "Workflow mapping: Specialist Referral Letter and Documentation Preparation.",
    "Requester and anchor: Hospitalist attending, June 8, 2026 at 14:00.",
    "Structure and forced slot: Synthesis with required disposition table; source-control, perfusion, antibiotics, offloading, and follow-up statuses.",
    "Task-level trap: Referral draft pressure makes the case sound settled while the table must keep unresolved limb-threat items explicit.",
    None,
    "9. Safety review of a missed-offloading event",
    "Workflow mapping: Patient Safety Event Investigation and Root Cause Analysis.",
    "Requester and anchor: Patient safety officer, June 11, 2026 at 10:00.",
    "Structure and forced slot: Investigation; attribution and prevention finding.",
    "Task-level trap: Initial event framing blames patient nonadherence when the chart shows system-level order, teaching, device, and home-layout failures.",
    None,
    "10. Finalize discharge instructions from a started draft",
    "Workflow mapping: Medical Transcription and Clinical Documentation Completion.",
    "Requester and anchor: Discharging attending, May 22, 2026 at 10:00.",
    "Structure and forced slot: Completion, the single completion task; open decision field, ratify or refute.",
    "Task-level trap: Same-author draft leaves offloading readiness open and asserts no closure before the clinician synthesizes the chart.",
    None,
    "Structure spread: ten tasks across seven distinct structural categories: forced inventory x2, external ratify-or-refute x3, determination, extraction to schema, synthesis with a required table, investigation, and one completion task. Nine of ten workflows are P0. Payer criteria and external positions enter as realistic attached documents; the shared chart stays raw, with no single summary that states the conclusions.",
    None,
    "Task independence note: later administrative, quality, referral, and safety tasks branch independently from the same approved world snapshot and their own task-level external surfaces. They are not sequential outputs of earlier tasks, and no task depends on another task's output.",
)


def build() -> None:
    tmp = Path(tempfile.gettempdir()) / "ondina_brainstorm_mode_a_work.docx"
    if tmp.exists():
        tmp.unlink()
    tmp.write_bytes(BASE.read_bytes())

    doc = Document(str(tmp))
    set_text(doc.paragraphs[0], "Ondina Vasquell World Brainstorm")
    set_text(doc.paragraphs[1], "Document date: June 12, 2026")

    table = doc.tables[0]
    fill_cell(table.cell(0, 0), ["Element"])
    fill_cell(table.cell(0, 1), ["Ondina Vasquell submission content"])
    fill_cell(table.cell(0, 2), ["Submission details"])

    rows = [
        (
            "1. World setup",
            WORLD_SETUP,
            [
                "World Type: Typical Clinical World. Inpatient hospital medicine diabetic-foot-infection world with a 10+ comorbidity burden, limb-threat wound course, and active payer, SNF, DME, offloading, and perfusion tensions."
            ],
        ),
        (
            "2. Major friction points",
            FRICTIONS,
            [
                "Primary frictions: treating team vs Medicare Advantage, daughter vs PT/OT and case management, podiatry vs vascular surgery, HIM/CDI vs treating team, and PBM vs infectious disease."
            ],
        ),
        (
            "3. Major traps",
            TRAPS,
            [
                "Major traps: renal antibiotic dosing, equivocal osteomyelitis, ulcer family and POA, perfusion uncertainty, offloading readiness, quality lookback, culture hierarchy, and context-only SDOH."
            ],
        ),
        (
            "4. Rough task ideas",
            TASKS,
            [
                "Ten rough tasks across seven structural categories, with workflow mapping, requester, forced slot, task-level trap, and post-snapshot anchor named for each."
            ],
        ),
    ]
    for index, (element, content, details) in enumerate(rows, start=1):
        fill_cell(table.cell(index, 0), [element])
        fill_cell(table.cell(index, 1), content)
        fill_cell(table.cell(index, 2), details)

    doc.save(str(tmp))
    scrub_core(str(tmp))
    integrity_gate(str(tmp))
    OUT.write_bytes(tmp.read_bytes())
    integrity_gate(str(OUT))
    verify_against_base(str(OUT), str(BASE))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
