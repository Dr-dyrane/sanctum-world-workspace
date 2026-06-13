#!/usr/bin/env python3
"""Build the Ondina Vasquell World Spec DOCX via Mode A clone of the approved KM spec.

Clones the KM spec base so styles.xml, theme, and numbering are byte-identical,
clears the body, and rebuilds it from the audited submission markdown. To keep the
Mode A fingerprint empty, the generated tables reproduce the base's exact color
vocabulary: fills {4472C4, D6E4F0, e8f5e9} and border colors {000000, 43a047,
4472C4, CCCCCC}, and nothing else.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
from tools.mode_a_clone import integrity_gate, scrub_core, verify_against_base  # noqa: E402

BASE = REPO / "worlds/korvin-merrow/file-review/pipeline-output/.meta/references/Alexander_World_Merrow_latest_6_4.docx"
SRC = REPO / "worlds/ondina-vasquell/submission/Ondina_Vasquell_World_Spec.md"
OUT = REPO / "worlds/ondina-vasquell/submission/Alexander_World_Vasquell_latest_6_13.docx"

HEADER_FILL = "4472C4"   # data-table header rows
META_FILL = "D6E4F0"     # the top identity/metadata table
CALLOUT_FILL = "e8f5e9"  # the green build-note callout
WHITE = "FFFFFF"


def clear_body(doc) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def shade(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def white_bold(cell) -> None:
    for p in cell.paragraphs:
        for run in p.runs:
            run.bold = True
            c = OxmlElement("w:color")
            c.set(qn("w:val"), WHITE)
            run._r.get_or_add_rPr().append(c)


def set_borders(table, outer: str, inner: str) -> None:
    borders = OxmlElement("w:tblBorders")
    spec = {"top": outer, "left": outer, "bottom": outer, "right": outer,
            "insideH": inner, "insideV": inner}
    for edge, color in spec.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def add_data_table(doc, rows, *, meta=False) -> None:
    if not rows:
        return
    ncol = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "TableNormal"
    if meta:
        # identity/metadata table: light-blue rows, black border (uses META_FILL + 000000)
        set_borders(t, outer="000000", inner="000000")
        for r in range(len(rows)):
            for c in range(ncol):
                t.cell(r, c).text = rows[r][c] if c < len(rows[r]) else ""
                shade(t.cell(r, c), META_FILL)
            t.cell(r, 0).paragraphs[0].runs and setattr(t.cell(r, 0).paragraphs[0].runs[0], "bold", True)
    else:
        # data tables: blue header row, white bold text, blue outer + grey inner borders
        set_borders(t, outer=HEADER_FILL, inner="CCCCCC")
        for r, row in enumerate(rows):
            for c in range(ncol):
                t.cell(r, c).text = row[c] if c < len(row) else ""
            if r == 0:
                for c in range(ncol):
                    shade(t.cell(r, c), HEADER_FILL)
                    white_bold(t.cell(r, c))


def add_callout(doc, text: str) -> None:
    """Green 1x1 callout box: e8f5e9 fill, 43a047 border (supplies both colors)."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "TableNormal"
    set_borders(t, outer="43a047", inner="43a047")
    t.cell(0, 0).text = text
    shade(t.cell(0, 0), CALLOUT_FILL)


def parse_md_table(block):
    rows = []
    for ln in block:
        stripped = ln.replace("|", "").replace("-", "").replace(":", "").strip()
        if stripped == "":
            continue
        rows.append([c.strip() for c in ln.strip().strip("|").split("|")])
    return rows


def build() -> None:
    md = SRC.read_text(encoding="utf-8")
    doc = Document(str(BASE))
    clear_body(doc)

    lines = md.split("\n")
    i, n = 0, len(lines)
    first_table_seen = False
    while i < n:
        ln = lines[i]
        if ln.startswith("|"):
            block = []
            while i < n and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            rows = parse_md_table(block)
            add_data_table(doc, rows, meta=not first_table_seen)
            first_table_seen = True
            continue
        if ln.startswith("# "):
            doc.add_paragraph(ln[2:].strip(), style="Heading 1")
        elif ln.startswith("## "):
            doc.add_paragraph(ln[3:].strip(), style="Heading 2")
        elif ln.startswith("### "):
            doc.add_paragraph(ln[4:].strip(), style="Heading 2")
        elif ln.startswith("Build note:"):
            add_callout(doc, ln.strip())
        elif ln.strip() == "":
            pass
        else:
            doc.add_paragraph(ln.strip(), style="normal")
        i += 1

    doc.save(str(OUT))
    scrub_core(str(OUT))
    print("integrity gate:", integrity_gate(str(OUT)))
    print("verify_against_base:", verify_against_base(str(OUT), str(BASE)))
    print("written:", OUT)


if __name__ == "__main__":
    build()
