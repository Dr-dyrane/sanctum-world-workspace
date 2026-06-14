#!/usr/bin/env python3
"""KM07 v2 -- nephrology referral letter completion, bone-health false closure, 05/26 anchor.
Built 2026-06-10 on the GENRE-TRUE base (primary_care_outpatient_baseline_summary world file:
clinic letterhead band, Author+NPI cells). Includes the post-build disposition-cell patch.
Outputs: platform/task7/current/nephrology_referral_letter_draft_05262026.docx + golden-KM07-v2.docx
NOTE: chrome em-dashes in the cloned band are swapped to hyphens (task-file precedent).
"""
import sys, os, tempfile, shutil, re, zipfile, hashlib

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
sys.path.insert(0, REPO_ROOT)
from tools.mode_a_clone import clone, set_text, insert_before, delete_paragraphs, scrub_core, integrity_gate
from docx import Document

REPO = os.path.join(REPO_ROOT, "worlds", "korvin-merrow")
BASE = os.path.join(REPO, "file-review", "upload", "filesystem", "primary_care_outpatient_baseline_summary_05182026.docx")
OUT  = os.path.join(REPO, "task-setup", "platform", "task7", "current")

REASON = ("Dear Dr. Solthar, thank you for seeing Mr. Merrow in follow-up after his 05/18 to "
    "05/24/2026 hospitalization for a suspected urinary-source infection with sepsis physiology "
    "and acute kidney injury on his baseline CKD stage 3. I am referring him for ongoing renal "
    "follow-up and to lead the staged reintroduction of the renal- and hemodynamic-sensitive "
    "medications held during the admission, in coordination with his cardiologist. None of the "
    "held agents has been restarted since discharge.")
COURSE = ("His outpatient baseline creatinine is 1.6 to 1.8 mg/dL with an eGFR in the 40 to 50 "
    "range. Admission creatinine peaked at 2.62 mg/dL with a multifactorial, predominantly "
    "prerenal injury. By discharge the creatinine had returned to 1.80 mg/dL, at the upper end "
    "of his baseline, and the AKI had resolved by laboratory criteria. Your inpatient service "
    "framed the medication restart as a staged, parameter-gated sequencing question guided by "
    "renal trajectory and blood pressure reserve rather than by any single favorable value.")
MEDS_DRAFT = ("Current medications: carvedilol 12.5 mg twice daily, aspirin 81 mg daily, "
    "atorvastatin 40 mg nightly, insulin glargine 18 units nightly, gabapentin 300 mg nightly, "
    "ferrous sulfate, calcium with vitamin D, pantoprazole, and alendronate 70 mg weekly "
    "(Sundays), which was continued through the admission; bone-health therapy is current with "
    "nothing outstanding on reconciliation. Held pending your guidance: sacubitril/valsartan "
    "24/26 mg twice daily, spironolactone 25 mg daily, furosemide 40 mg daily, empagliflozin "
    "10 mg daily, metformin ER 500 mg twice daily. Prednisone continues on its outpatient taper.")
MEDS_GOLD = ("Current medications: carvedilol 12.5 mg twice daily, aspirin 81 mg daily, "
    "atorvastatin 40 mg nightly, insulin glargine 18 units nightly, gabapentin 300 mg nightly, "
    "ferrous sulfate, calcium with vitamin D, and pantoprazole. Of note, alendronate 70 mg "
    "weekly (Sundays) was not administered during the admission and remains an open "
    "reconciliation item on the discharge list; given the recent AKI on CKD stage 3 and his "
    "chronic steroid exposure, I would ask you to confirm renal trajectory at your visit before "
    "it is resumed, and bone-health management otherwise remains an open outpatient item. Held "
    "pending your guidance: sacubitril/valsartan 24/26 mg twice daily, spironolactone 25 mg "
    "daily, furosemide 40 mg daily, empagliflozin 10 mg daily, metformin ER 500 mg twice "
    "daily. Prednisone continues on its outpatient taper.")
RESTART_GOLD = ("I want to be clear that laboratory resolution is not the same as restored "
    "physiologic reserve, and it does not by itself establish readiness to resume the held "
    "cardiorenal and diabetes agents. Your service framed the restart as a staged, "
    "parameter-gated sequencing problem, coordinated with cardiology rather than undertaken "
    "simultaneously. I am referring that open sequencing decision to you to drive: which agent "
    "first, against which parameters, and on what monitoring schedule. None of the held agents "
    "has been restarted since discharge.")
REQ_DRAFT = ("He would benefit from being seen within one to two weeks of discharge as your "
    "service recommended. Please advise on restart sequencing, monitoring intervals, and any "
    "interim laboratory checks you would like before his visit.")
REQ_GOLD = ("I would be grateful if he could be seen within one to two weeks of discharge as "
    "your service recommended, sooner if any restart is initiated. His wife oversees "
    "medications at home and he is ambulating with a rolling walker; please advise on any "
    "interim laboratory checks you would like before the visit.")

def build(which):
    tmp = os.path.join(tempfile.gettempdir(), f"km07v2_{which}.docx")
    d = clone(BASE, tmp)
    paras = d.paragraphs
    set_text(paras[2], "NEPHROLOGY REFERRAL LETTER" + (" - DRAFT" if which == "draft" else ""))
    set_text(paras[3], "Author: Talia Quenor, MD - Primary Care  |  Mercy Vale Primary Care, Harbor Crest Clinic  |  "
             "05/26/2026  |  Status: " + ("Draft for finalization" if which == "draft" else "Signed"))
    heading_like = paras[4]
    body_like    = paras[6]
    sig_anchor   = paras[44]
    sig_name     = paras[45]
    sig_clinic   = paras[46]
    delete_paragraphs(list(paras[4:44]))
    body = [(True, "REASON FOR REFERRAL"), (False, REASON),
            (True, "HOSPITAL COURSE"), (False, COURSE),
            (True, "MEDICATIONS AND INTERVAL STATUS"),
            (False, MEDS_DRAFT if which == "draft" else MEDS_GOLD)]
    if which == "golden":
        body += [(True, "RESTART SEQUENCING, THE OPEN QUESTION"), (False, RESTART_GOLD)]
    body += [(True, "REQUESTS AND FOLLOW-UP"),
             (False, REQ_DRAFT if which == "draft" else REQ_GOLD),
             (False, "To finalize: complete and send." if which == "draft" else
                     "Finalized and sent with the post-discharge records; alendronate reconciliation left open pending your renal reassessment.")]
    for is_h, t in body:
        insert_before(sig_anchor, heading_like if is_h else body_like, t)
    if which == "draft":
        set_text(sig_anchor, "Draft started 05/26/2026 - complete and send")
        set_text(sig_name, "For: Talia Quenor, MD - Primary Care")
    else:
        set_text(sig_anchor, "Electronically signed 05/26/2026")
        set_text(sig_name, "Talia Quenor, MD - Primary Care")
    set_text(sig_clinic, "Mercy Vale Primary Care - Harbor Crest Clinic  |  NPI 1962718304  |  (412) 555-0410")
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                tx = c.text
                for p in c.paragraphs:
                    if '05/18/2026' in p.text and 'Document date' in tx:
                        set_text(p, '05/26/2026')
                    elif p.text.strip() == 'Primary care outpatient baseline summary':
                        set_text(p, 'Nephrology referral letter' + (' - draft' if which == 'draft' else ''))
                    elif 'Forwarded to inpatient team' in p.text:
                        set_text(p, 'To: Iven Solthar, MD - Nephrology, Mercy Vale Regional Medical Center')
                    elif '—' in p.text or '–' in p.text:
                        set_text(p, p.text.replace('—', '-').replace('–', '-'))
    for p in d.paragraphs:
        if '—' in p.text or '–' in p.text:
            set_text(p, p.text.replace('—', '-').replace('–', '-'))
    d.save(tmp); scrub_core(tmp); integrity_gate(tmp)
    s_new = hashlib.sha256(zipfile.ZipFile(tmp).read('word/styles.xml')).hexdigest()
    s_base = hashlib.sha256(zipfile.ZipFile(BASE).read('word/styles.xml')).hexdigest()
    txt = '\n'.join(p.text for p in Document(tmp).paragraphs)
    for t in Document(tmp).tables:
        for r in t.rows:
            for c in r.cells: txt += '\n' + c.text
    print(f"=== {which.upper()} ===")
    print('  styles byte-identical:', s_new == s_base)
    print('  em/en/arrow count:', sum(txt.count(ch) for ch in ['—', '–', '→']))
    print('  synthetic token:', 'Synthetic' in txt)
    fname = "nephrology_referral_letter_draft_05262026.docx" if which == "draft" else "golden-KM07-v2.docx"
    shutil.copy(tmp, os.path.join(OUT, fname)); print('  Saved:', fname)

if __name__ == "__main__":
    build("draft"); build("golden")
