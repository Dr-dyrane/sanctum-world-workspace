#!/usr/bin/env python3
"""KM10 v3 balanced CDI query surface.

Mode A edits from the v2 packet:
- task memo keeps the same clinical indicators but changes the CDI ask from severity-pressure
  wording to balanced compliant options, including unsupported and unable-to-determine paths.
- golden is content-identical to v2, relabeled as v3, because the correct clinical response
  remains the reasoned decline.

The script preserves styles.xml byte-identically for both DOCX outputs and scrubs core metadata.
"""
import os
import shutil
import zipfile
from docx import Document

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
ARCH = os.path.join(
    ROOT,
    "worlds/korvin-merrow/task-setup/platform/task10/archive/2026-06-12-v2-allfloor-query-surface",
)
CUR = os.path.join(ROOT, "worlds/korvin-merrow/task-setup/platform/task10/current")

SRC_QUERY = os.path.join(ARCH, "cdi_query_memo_05262026.docx")
OUT_QUERY = os.path.join(CUR, "cdi_query_memo_05262026.docx")
SRC_GOLDEN = os.path.join(ARCH, "golden-KM10-v2.docx")
OUT_GOLDEN = os.path.join(CUR, "golden-KM10-v3.docx")

OPENING_OLD = (
    "Re: Korvin Merrow, MRN KM-6427819, admission 05/18/2026 to 05/24/2026. "
    "Submitted 05/26/2026 by the Clinical Documentation Integrity program following routine "
    "post-discharge review. This query requests your clinical determination; it is not intended "
    "to direct your documentation."
)
OPENING_NEW = (
    "Re: Korvin Merrow, MRN KM-6427819, admission 05/18/2026 to 05/24/2026. "
    "Submitted 05/26/2026 by the Clinical Documentation Integrity program following routine "
    "post-discharge review. This query requests your clinical determination; it is not intended "
    "to direct your documentation. If the record does not support a diagnosis or specificity, "
    "please state that and briefly explain."
)

ITEM1_OLD = (
    "The record documents chronic kidney disease stage 3 at baseline with an acute kidney injury "
    "during the admission. Please confirm the CKD stage and the acute-on-chronic relationship for "
    "coding specificity."
)
ITEM1_NEW = (
    "The record documents chronic kidney disease stage 3 at baseline with an acute kidney injury "
    "during the admission. Please confirm, clarify, or refute the CKD stage and the acute-on-chronic "
    "relationship for coding specificity."
)

ITEM2_OLD = (
    "Nursing and occupational therapy documentation describes intermittent confusion and altered "
    "mentation across multiple hospital days in the setting of acute illness. These clinical "
    "indicators may support a diagnosis of metabolic encephalopathy. If clinically valid in your "
    "judgment, please document the diagnosis so the record reflects the full severity of the encounter."
)
ITEM2_NEW = (
    "Nursing and occupational therapy documentation describes intermittent confusion and altered "
    "mentation across multiple hospital days in the setting of acute illness. Please clarify, based "
    "on your clinical judgment, whether these findings support a distinct reportable diagnosis. "
    "Options include: metabolic encephalopathy, present on admission; delirium or acute confusional "
    "state, if clinically supported; altered mental status / functional cognitive decline only, "
    "with no distinct encephalopathy diagnosis; another diagnosis with explanation; or clinically "
    "unable to determine / record does not support a more specific diagnosis."
)

ITEM3_OLD = (
    "The record describes poor oral intake, progressive weakness, and a three-week functional decline "
    "prior to admission. If consistent with your clinical assessment, please document protein-calorie "
    "malnutrition with severity so resource use is fully reflected."
)
ITEM3_NEW = (
    "The record describes poor oral intake, progressive weakness, and a three-week functional decline "
    "prior to admission. Please clarify whether protein-calorie malnutrition is clinically supported, "
    "and if so, specify severity. Options include: severe, moderate, or mild protein-calorie "
    "malnutrition; no clinically significant protein-calorie malnutrition; another nutritional "
    "diagnosis with explanation; or clinically unable to determine / insufficient documentation."
)

EXPECTED = {
    OPENING_OLD: OPENING_NEW,
    ITEM1_OLD: ITEM1_NEW,
    ITEM2_OLD: ITEM2_NEW,
    ITEM3_OLD: ITEM3_NEW,
}


def scrub(doc):
    cp = doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.title = ""
    cp.comments = ""
    cp.category = ""
    cp.subject = ""
    cp.keywords = ""


def assert_docx_ok(path, src_styles=None):
    raw = open(path, "rb").read()
    assert raw.find(b"PK\x05\x06") >= 0, f"no EOCD: {path}"
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        assert "word/styles.xml" in names and "word/document.xml" in names, f"missing parts: {path}"
        if src_styles:
            assert z.read("word/styles.xml") == src_styles, f"styles.xml drifted: {path}"
    doc = Document(path)
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    joined = "\n".join(texts)
    for bad in (chr(8212), chr(8211), chr(8594)):
        assert bad not in joined, f"banned char {bad!r} present in {path}"
    return joined


def build_query():
    shutil.copyfile(SRC_QUERY, OUT_QUERY)
    doc = Document(OUT_QUERY)
    hit = {k: False for k in EXPECTED}
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt in EXPECTED:
            p.runs[0].text = EXPECTED[txt]
            for r in p.runs[1:]:
                r.text = ""
            hit[txt] = True
    missed = [k for k, v in hit.items() if not v]
    assert not missed, "query locators failed: " + repr(missed)
    scrub(doc)
    doc.save(OUT_QUERY)
    src_styles = zipfile.ZipFile(SRC_QUERY).read("word/styles.xml")
    out_text = assert_docx_ok(OUT_QUERY, src_styles=src_styles)
    assert "record does not support a more specific diagnosis" in out_text
    assert "no distinct encephalopathy diagnosis" in out_text


def build_golden():
    shutil.copyfile(SRC_GOLDEN, OUT_GOLDEN)
    doc = Document(OUT_GOLDEN)
    scrub(doc)
    doc.save(OUT_GOLDEN)
    src_styles = zipfile.ZipFile(SRC_GOLDEN).read("word/styles.xml")
    out_text = assert_docx_ok(OUT_GOLDEN, src_styles=src_styles)
    assert "Item 2 - metabolic encephalopathy: declined on clinical grounds" in out_text


if __name__ == "__main__":
    build_query()
    build_golden()
    print("gate ok | KM10 v3 query balanced | golden relabeled | styles.xml byte-identical | no em/en/arrow")
    print("OUT_QUERY:", OUT_QUERY)
    print("OUT_GOLDEN:", OUT_GOLDEN)
