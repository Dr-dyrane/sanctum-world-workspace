#!/usr/bin/env python3
"""Build the Marva Lydell Brainstorm DOCX via Mode A clone."""

from __future__ import annotations

import copy
import re
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))

from tools.mode_a_clone import integrity_gate, scrub_core, set_text, verify_against_base


BASE = REPO / "worlds/korvin-merrow/submission/Korvin_Merrow_Brainstorm.docx"
SOURCE = REPO / "worlds/marva-lydell/submission/Marva_Lydell_Brainstorm.md"
OUT = REPO / "worlds/marva-lydell/submission/Marva_Lydell_Brainstorm.docx"
BANNED_RE = re.compile("[\u2013\u2014\u2190-\u21ff]")


def remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._p.getparent().remove(paragraph._p)


def fill_cell(cell, lines: list[str]) -> None:
    """Replace a table cell with paragraph clones from its first paragraph."""
    if not lines:
        lines = [""]
    proto = copy.deepcopy(cell.paragraphs[0]._p)
    for paragraph in list(cell.paragraphs):
        remove_paragraph(paragraph)
    for text in lines:
        new_el = copy.deepcopy(proto)
        new_p = Paragraph(new_el, cell)
        set_text(new_p, text)
        cell._tc.append(new_el)


def extract_section(text: str, heading: str) -> list[str]:
    marker = f"## {heading}"
    start = text.index(marker) + len(marker)
    rest = text[start:].splitlines()
    lines: list[str] = []
    for line in rest:
        if line.startswith("## "):
            break
        lines.append(line.rstrip())
    while lines and not lines[0].strip():
        lines.pop(0)
    while lines and not lines[-1].strip():
        lines.pop()
    return lines


def extract_preamble(text: str) -> list[str]:
    lines = text.splitlines()
    out: list[str] = []
    seen_title = False
    for line in lines:
        if line.startswith("# "):
            seen_title = True
            continue
        if seen_title and line.startswith("## "):
            break
        if seen_title:
            out.append(line.rstrip())
    while out and not out[0].strip():
        out.pop(0)
    while out and not out[-1].strip():
        out.pop()
    return out


def table_cells(line: str) -> list[str] | None:
    stripped = line.strip()
    if not stripped.startswith("|") or not stripped.endswith("|"):
        return None
    cells = [cell.strip() for cell in stripped.strip("|").split("|")]
    if all(cell and set(cell) <= {"-", ":", " "} for cell in cells):
        return []
    if cells and cells[0] in {"Field", "Friction", "Trap", "ID", "#"}:
        return []
    return cells


def convert_markdown(lines: list[str]) -> list[str]:
    out: list[str] = []
    for raw in lines:
        line = raw.strip()
        if not line:
            if out and out[-1] != "":
                out.append("")
            continue
        if line.startswith("#"):
            continue

        cells = table_cells(line)
        if cells is not None:
            if not cells:
                continue
            if len(cells) == 2:
                out.append(f"{cells[0]}: {cells[1]}")
            elif len(cells) == 3:
                out.append(f"{cells[0]}: {cells[1]}")
                out.append(f"Supporting record: {cells[2]}")
            elif len(cells) >= 8:
                out.append(f"{cells[0]}. {cells[2]}")
                out.append(f"Requester: {cells[1]}")
                out.append(f"Workflow: {cells[3]}")
                out.append(f"Structure: {cells[4]}")
                out.append(f"Required decision: {cells[5]}")
                out.append(f"Clinical trap: {cells[6]}")
                out.append(f"Anchor: {cells[7]}")
            else:
                out.append(" | ".join(cells))
            continue

        if line.startswith("- "):
            out.append(line)
        else:
            out.append(line)
    while out and out[-1] == "":
        out.pop()
    return out


def build() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    if BANNED_RE.search(text):
        raise ValueError("Source contains banned dash or arrow characters")

    tmp = Path(tempfile.gettempdir()) / "marva_brainstorm_mode_a_work.docx"
    if tmp.exists():
        tmp.unlink()
    tmp.write_bytes(BASE.read_bytes())

    doc = Document(str(tmp))
    set_text(doc.paragraphs[0], "Marva Lydell World Brainstorm")
    set_text(doc.paragraphs[1], "Document date: July 18, 2025")

    concept = convert_markdown(extract_preamble(text))
    world_setup = concept + [""] + convert_markdown(extract_section(text, "1. World setup"))
    frictions = convert_markdown(extract_section(text, "2. Major friction points"))
    traps = convert_markdown(extract_section(text, "3. Major traps"))
    tasks = convert_markdown(extract_section(text, "4. Rough task ideas"))

    table = doc.tables[0]
    fill_cell(table.cell(0, 0), ["Element"])
    fill_cell(table.cell(0, 1), ["Marva Lydell submission content"])
    fill_cell(table.cell(0, 2), ["Submission details"])

    rows = [
        (
            "1. World setup",
            world_setup,
            [
                "World Type: Typical Clinical World. Cardiorenal respiratory transition world with readiness, oxygen, DME, renal medication, anticoagulation, payer, and home-support frictions."
            ],
        ),
        (
            "2. Major friction points",
            frictions,
            [
                "Primary frictions: treating team vs payer, cardiology vs nephrology, respiratory therapy and pulmonary vs DME execution, patient and family preference vs functional safety, and pharmacy or SNF intake vs the inpatient chart."
            ],
        ),
        (
            "3. Major traps",
            traps,
            [
                "Major traps: resting oxygen over-closure, missing equipment, volume-status over-closure, renal medication drift, anticoagulation source hierarchy, home support inflation, respiratory control over-closure, denial authority bias, and readmission attribution shortcut."
            ],
        ),
        (
            "4. Rough task ideas",
            tasks,
            [
                "Ten rough tasks across claims, medication reconciliation, utilization review, documentation completion, coordination, referral, safety, quality abstraction, and CDI workflows."
            ],
        ),
    ]
    for index, (element, content, details) in enumerate(rows, start=1):
        fill_cell(table.cell(index, 0), [element])
        fill_cell(table.cell(index, 1), content)
        fill_cell(table.cell(index, 2), details)

    doc.save(str(tmp))
    scrub_core(str(tmp))
    integrity_gate(str(tmp))
    OUT.write_bytes(tmp.read_bytes())
    integrity_gate(str(OUT))
    verify_against_base(str(OUT), str(BASE))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
