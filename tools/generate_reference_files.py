#!/usr/bin/env python3
"""Reference-file generator: locked FI/EW markdown -> styled clinical DOCX.

Reusable across worlds. Design spec: worlds/<world>/reference-file-design/epic-note-design-system.md
Approved visual sample: FI-W01 artisan build (reference-file-design/samples/).

Per-file pipeline: parse metadata header + ## sections -> render with chrome:
synthetic banner (top+bottom, centered), facility masthead, patient storyboard,
title+filing line, ruled section headers, bullets, md tables (blue header rows),
pull-quotes for '...states: "..."' lines, signature block, running header/footer.

Guardrails baked in: renders locked content only; strips canon-guard lines
(post-world / golden / grader mentions); NO em dashes/arrows in chrome;
integrity gate after save (opens + EOCD). Translate IDs via NEW_IDS before
rendering so files match the spec's file-plan convention.

Usage: python tools/generate_reference_files.py  (edit JOBS/IDS/FN at bottom)
"""
import re, os, glob
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE=RGBColor(0x44,0x72,0xC4); NAVY=RGBColor(0x1F,0x38,0x64); INK=RGBColor(0x23,0x28,0x30)
GRAY=RGBColor(0x5E,0x66,0x70); WHITE=RGBColor(0xFF,0xFF,0xFF); AMBER=RGBColor(0x9A,0x60,0x00)
B_HEX="4472C4"; CARD="EDF2FA"; CREAM="FFF8E1"; RULE="C9D4EA"
FACILITY="MERCY VALE REGIONAL MEDICAL CENTER"
PATIENT=("KORVIN MERROW","62 y  |  Male  |  DOB 02/18/1964","MRN KM-6427819","Lisinopril (cough)")
BANNER="SYNTHETIC TRAINING DOCUMENT  |  FICTIONAL PATIENT  |  NOT A REAL MEDICAL RECORD"
STRIP=re.compile(r'no discharge outcome|post-world|task prompt|golden response|grader guidance|expected output',re.I)
DEPT=[("triage","Emergency Department"),("ed ","Emergency Department"),("hospitalist","Hospital Medicine"),
 ("nephrology","Nephrology"),("cardiology","Cardiology"),("endocrinology","Endocrinology"),
 ("rheumatology","Outpatient Rheumatology"),("primary care","Primary Care"),("pharmacy","Pharmacy"),
 ("refill","Pharmacy"),("medication administration","Pharmacy"),("nursing","Nursing"),
 ("physical therapy","Rehabilitation Services"),("occupational","Rehabilitation Services"),
 ("case management","Care Management"),("family communication","Care Management"),
 ("discharge-facing","Hospital Medicine"),("trend summary","Clinical Data / Laboratory"),
 ("problem list","Medical Records"),("pci","Cardiology Records"),("sleep","Sleep Medicine Records"),
 ("home support","Care Management"),("request context","Care Coordination"),("addendum","Pharmacy")]
NEW_IDS={}  # e.g. {"FI-W01":"EW1", ...} - fill per world's spec convention

def dept_for(ft):
    f=ft.lower()
    for k,v in DEPT:
        if k in f: return v
    return "Medical Records"

def repl_ids(s):
    for old,new in sorted(NEW_IDS.items(),key=lambda x:-len(x[0])): s=s.replace(old,new)
    return s

def parse(path):
    meta={}; body=[]; insec=False
    for ln in open(path,encoding='utf-8').read().split('\n'):
        ln=repl_ids(ln)
        if ln.startswith('## '): insec=True; body.append(('h',ln[3:].strip())); continue
        if not insec:
            m=re.match(r'([A-Za-z /()#]+):\s*(.+)$',ln)
            if m: meta[m.group(1).strip()]=m.group(2).strip()
            continue
        if ln.startswith('| '): body.append(('t',ln)); continue
        if re.match(r'^[-*] ',ln): body.append(('b',ln[2:].strip())); continue
        if ln.strip():
            if STRIP.search(ln) and len(ln)<200: continue
            body.append(('p',ln.strip()))
    return meta,body

def build(src,out):
    meta,body=parse(src)
    ftype=meta.get('File Type','Clinical document'); date=meta.get('Approximate Date / Anchor','')
    author=meta.get('Author / Source','Clinical staff'); dept=dept_for(ftype)
    d=Document(); sec=d.sections[0]
    sec.top_margin=Inches(0.5); sec.bottom_margin=Inches(0.55); sec.left_margin=Inches(0.65); sec.right_margin=Inches(0.65)
    n=d.styles['Normal']; n.font.name='Arial'; n.font.size=Pt(9.5); n.font.color.rgb=INK
    n.paragraph_format.space_after=Pt(4); n.paragraph_format.line_spacing=1.08
    def shade(c,f):
        p=c._tc.get_or_add_tcPr(); s=OxmlElement('w:shd'); s.set(qn('w:val'),'clear'); s.set(qn('w:fill'),f); p.append(s)
    def borders(t,none=False):
        pr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
        for e in ('top','left','bottom','right','insideH','insideV'):
            el=OxmlElement('w:'+e)
            if none: el.set(qn('w:val'),'none')
            else: el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),RULE)
            b.append(el)
        pr.append(b)
    def run(p,t,sz=9.5,c=INK,b=False,i=False):
        r=p.add_run(t); r.font.name='Arial'; r.font.size=Pt(sz); r.font.color.rgb=c; r.font.bold=b; r.font.italic=i
    def rich(p,text,sz=9.5):
        for tok in re.split(r'(\*\*[^*]+\*\*)',text):
            if not tok: continue
            if tok.startswith('**'): run(p,tok[2:-2],sz,b=True)
            else: run(p,tok,sz)
    def band():
        t=d.add_table(rows=1,cols=1); borders(t,none=True); c=t.rows[0].cells[0]; shade(c,CREAM)
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run(p,BANNER,8,AMBER,b=True)
    def header(t):
        p=d.add_paragraph(); p.paragraph_format.space_before=Pt(9); p.paragraph_format.space_after=Pt(3)
        run(p,t.upper(),10,BLUE,b=True)
        pr=p._p.get_or_add_pPr(); pb=OxmlElement('w:pBdr'); bo=OxmlElement('w:bottom')
        bo.set(qn('w:val'),'single'); bo.set(qn('w:sz'),'6'); bo.set(qn('w:color'),RULE); bo.set(qn('w:space'),'2')
        pb.append(bo); pr.append(pb)
    band()
    m=d.add_table(rows=1,cols=2); borders(m,none=True)
    run(m.rows[0].cells[0].paragraphs[0],FACILITY,11,NAVY,b=True)
    run(m.rows[0].cells[0].add_paragraph(),dept,8,GRAY)
    pr=m.rows[0].cells[1].paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.RIGHT; run(pr,"Confidential",8.5,BLUE,b=True)
    hr=d.add_table(rows=1,cols=1); borders(hr,none=True); shade(hr.rows[0].cells[0],B_HEX)
    tr=hr.rows[0]._tr.get_or_add_trPr(); h=OxmlElement('w:trHeight'); h.set(qn('w:val'),'40'); tr.append(h)
    sb=d.add_table(rows=2,cols=4); borders(sb)
    top=sb.rows[0]; mc=top.cells[0].merge(top.cells[2]); shade(mc,B_HEX)
    p=mc.paragraphs[0]; run(p,"  "+PATIENT[0],13,WHITE,b=True); run(p,"    "+PATIENT[1],9,WHITE)
    c3=top.cells[3]; shade(c3,B_HEX); p=c3.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; run(p,PATIENT[2]+"  ",10,WHITE,b=True)
    for j,(k,v) in enumerate([("Date / Anchor",date),("Source",author),("Allergies",PATIENT[3]),("Document",ftype)]):
        cell=sb.rows[1].cells[j]; shade(cell,CARD)
        p=cell.paragraphs[0]; run(p,k+"\n",7.4,GRAY,b=True); run(p,v[:90],8.2)
    d.add_paragraph()
    p=d.add_paragraph(); run(p,ftype.upper(),14,NAVY,b=True)
    p=d.add_paragraph(); run(p,f"Author: {author}  |  Department: {dept}  |  {date}  |  Status: Signed",8,GRAY)
    i=0
    while i<len(body):
        kind,val=body[i]
        if kind=='h': header(val)
        elif kind=='b':
            p=d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(1.5); rich(p,val)
        elif kind=='t':
            rows=[]
            while i<len(body) and body[i][0]=='t':
                cells=[c.strip() for c in body[i][1].strip().strip('|').split('|')]
                if not re.match(r'^:?-{2,}',cells[0]): rows.append(cells)
                i+=1
            i-=1
            if rows:
                cols=max(len(r) for r in rows)
                tb=d.add_table(rows=len(rows),cols=cols); borders(tb)
                for ri,r in enumerate(rows):
                    for ci in range(cols):
                        c=tb.rows[ri].cells[ci]; pp=c.paragraphs[0]
                        v=r[ci] if ci<len(r) else ''
                        if ri==0: shade(c,B_HEX); run(pp,v,8,WHITE,b=True)
                        else:
                            if ri%2==0: shade(c,CARD)
                            run(pp,v,8)
                d.add_paragraph()
        else:
            mq=re.match(r'(.*?(?:states|said|stated)[^:]*:)\s*"(.+)"$',val)
            if mq:
                p=d.add_paragraph(); rich(p,mq.group(1),9)
                q=d.add_paragraph(); q.paragraph_format.left_indent=Inches(0.35)
                qr=q._p.get_or_add_pPr(); pb=OxmlElement('w:pBdr'); le=OxmlElement('w:left')
                le.set(qn('w:val'),'single'); le.set(qn('w:sz'),'18'); le.set(qn('w:color'),B_HEX); le.set(qn('w:space'),'8')
                pb.append(le); qr.append(pb)
                run(q,'"'+mq.group(2)+'"',10.5,NAVY,i=True)
            else:
                p=d.add_paragraph(); rich(p,val)
        i+=1
    sg=d.add_paragraph(); sg.paragraph_format.space_before=Pt(10)
    pr3=sg._p.get_or_add_pPr(); pb=OxmlElement('w:pBdr'); bo=OxmlElement('w:top')
    bo.set(qn('w:val'),'single'); bo.set(qn('w:sz'),'4'); bo.set(qn('w:color'),RULE); bo.set(qn('w:space'),'4')
    pb.append(bo); pr3.append(pb)
    run(sg,f"Electronically signed by {author}  |  {dept}",9,GRAY,i=True)
    band()
    run(sec.header.paragraphs[0],f"{PATIENT[0].title()}  |  {PATIENT[2]}",7.5,GRAY)
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run(fp,f"{FACILITY.title()}  |  {ftype}  |  Synthetic training document",7.5,GRAY)
    d.save(out)
    raw=open(out,'rb').read(); assert raw.find(b'PK\x05\x06')>=0, f"corrupt save: {out}"
    Document(out)

if __name__=='__main__':
    # EXAMPLE (Korvin): map sources -> date-stamped outputs, set NEW_IDS, then:
    # for src,out in JOBS: build(src,out)
    print("Configure JOBS/NEW_IDS at bottom of file, then run. See world-pipeline-playbook.md.")
