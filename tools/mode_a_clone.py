#!/usr/bin/env python3
"""Mode A clone helper - the canonical build method for Korvin task docx (one Epic UI).

WHY: we mimic ONE Epic UI across every artifact. New task files and goldens are built
by CLONING a proven, approved artifact of the same UI and swapping ONLY the content,
keeping styles.xml byte-identical. Never build from blank Document(). Do NOT use
tools/generate_reference_files.py for task artifacts (regressed: 2-row band, injects a
"Synthetic training document" footer). See docs/reasoning-discipline.md (MODE A section)
and docs/docx-generation-method.md.

CANONICAL BASES (proven, approved, Arial, 3-row Epic band, clean footer, no banner):
  TASK FILES: worlds/korvin-merrow/task-setup/platform/task2/current/discharge_summary_draft_incomplete_05242026.docx
  GOLDENS:    worlds/korvin-merrow/task-setup/platform/task2/current/golden-KM02-v5.docx

This module gives the primitives (clone, edit-in-place, deepcopy-insert, band/footer
edit) and a FINGERPRINT-DIFF verifier. "Matches" means an EMPTY fingerprint diff, not
"looks close." Build per artifact in a short script that imports these; see the
__main__ example.
"""
import copy, zipfile, re, collections, hashlib, shutil
import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def clone(base_path, out_path):
    """Copy the approved base docx to out_path (keeps styles.xml/theme byte-identical)."""
    shutil.copy(base_path, out_path)
    return docx.Document(out_path)


def set_text(p, text):
    """Replace a paragraph's text while keeping the first run's formatting."""
    if not p.runs:
        p.add_run(text); return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def insert_before(anchor_p, like_p, text):
    """Insert a new paragraph (cloned from like_p's formatting) before anchor_p."""
    newp = copy.deepcopy(like_p._p)
    np = Paragraph(newp, like_p._parent)
    runs = np.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    anchor_p._p.addprevious(newp)
    return np


def delete_paragraphs(plist):
    for p in plist:
        p._p.getparent().remove(p._p)


def edit_band_cell(d, label_prefix, new_value_for_old, old_token):
    """In the identity band, replace old_token with new value in the cell whose text
    starts with label_prefix (e.g. 'Service', 'Document')."""
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                if c.text.strip().startswith(label_prefix) and old_token in c.text:
                    for p in c.paragraphs:
                        for run in p.runs:
                            if old_token in run.text:
                                run.text = run.text.replace(old_token, new_value_for_old)


def edit_footer(d, old_token, new_token):
    for s in d.sections:
        for fp in s.footer.paragraphs:
            for run in fp.runs:
                if old_token in run.text:
                    run.text = run.text.replace(old_token, new_token)


def integrity_gate(path):
    raw = open(path, "rb").read()
    assert raw.find(b"PK\x05\x06") >= 0, f"no EOCD (truncated): {path}"
    names = zipfile.ZipFile(path).namelist()
    assert "word/styles.xml" in names and "word/document.xml" in names
    docx.Document(path)
    return True


def fingerprint(path):
    doc = zipfile.ZipFile(path).read("word/document.xml").decode("utf8")
    return {
        "sizes": dict(collections.Counter(re.findall(r'<w:sz w:val="(\d+)"', doc))),
        "colors": set(re.findall(r'<w:color w:val="([^"]+)"', doc)),
        "fills": set(re.findall(r'w:fill="([^"]+)"', doc)),
        "borders": set(re.findall(r'<w:(?:top|bottom|left|right|insideH|insideV)\b[^/>]*w:color="([^"]+)"', doc)),
        "styles_md5": hashlib.md5(zipfile.ZipFile(path).read("word/styles.xml")).hexdigest(),
        "emdash": doc.count("—"), "endash": doc.count("–"), "arrow": doc.count("→"),
        "synthetic": "Synthetic" in doc, "banner": "SYNTHETIC TRAINING" in doc,
    }


def _core_clean(path):
    try:
        c = zipfile.ZipFile(path).read("docProps/core.xml").decode()
    except KeyError:
        return True
    creator_clean = "<dc:creator></dc:creator>" in c or "<dc:creator/>" in c
    description_clean = "<dc:description></dc:description>" in c or "<dc:description/>" in c
    modified_by_clean = "<cp:lastModifiedBy></cp:lastModifiedBy>" in c or "<cp:lastModifiedBy/>" in c
    return creator_clean and description_clean and modified_by_clean and "python-docx" not in c


def verify_against_base(out_path, base_path):
    """The definition of 'matches the template': empty fingerprint diff on chrome."""
    o, b = fingerprint(out_path), fingerprint(base_path)
    checks = {
        "styles.xml byte-identical": o["styles_md5"] == b["styles_md5"],
        "fills identical": o["fills"] == b["fills"],
        "border colors identical": o["borders"] == b["borders"],
        "palette subset of base": o["colors"] <= b["colors"],
        "no em/en dash/arrow": o["emdash"] == 0 and o["endash"] == 0 and o["arrow"] == 0,
        "no synthetic token": not o["synthetic"],
        "no banner": not o["banner"],
        "core metadata scrubbed": _core_clean(out_path),
    }
    for k, v in checks.items():
        print(f"  [{'OK ' if v else 'FAIL'}] {k}")
    assert all(checks.values()), "fingerprint diff is NOT empty - fix before shipping"
    print("  => fingerprint diff empty: MATCHES the approved base.")
    return True



def scrub_core(path):
    """Empty the python-docx fingerprint (dc:creator, dc:description, lastModifiedBy)
    to match the engineering-pipeline world files (which are scrubbed blank). python-docx
    writes creator='python-docx' / description='generated by python-docx' on save; that is
    a 'No Embedded Metadata Leaks' AutoQC fail. Run this AFTER the final save of any
    python-docx-built artifact, before staging."""
    import os
    tmp = path + ".tmp"
    zin = zipfile.ZipFile(path, "r"); zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "docProps/core.xml":
            c = data.decode("utf-8")
            c = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator></dc:creator>", c)
            c = re.sub(r"<dc:description>.*?</dc:description>", "<dc:description></dc:description>", c)
            c = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", "<cp:lastModifiedBy></cp:lastModifiedBy>", c)
            data = c.encode("utf-8")
        zout.writestr(item, data)
    zin.close(); zout.close(); os.replace(tmp, path)


_SYN = ("synthetic training document", "synthetic training", "synthetic")
_TOOL = ("python-docx", "aspose")


def scrub_all_metadata(path):
    """Phase 3 world/task files: scrub EVERY metadata surface, not just core.xml.
    Empties core.xml (creator/description/lastModifiedBy, python-docx), neutralizes
    app.xml value fields (Application/Company/Template/Manager), drops custom.xml,
    and blanks any <w:t> carrying a SYNTHETIC marker in document/headers/footers.
    The footer/body synthetic text is best removed at build time by clearing the
    footer; this is the belt-and-braces pass. Run AFTER final save, before staging,
    then assert verify_no_synthetic(path)."""
    import os
    tmp = path + ".tmp"
    zin = zipfile.ZipFile(path, "r"); zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        name = item.filename
        if name == "docProps/custom.xml":
            continue  # drop custom properties entirely
        data = zin.read(name)
        if name == "docProps/core.xml":
            c = data.decode("utf-8")
            c = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator></dc:creator>", c)
            c = re.sub(r"<dc:description>.*?</dc:description>", "<dc:description></dc:description>", c)
            c = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", "<cp:lastModifiedBy></cp:lastModifiedBy>", c)
            data = c.encode("utf-8")
        elif name == "docProps/app.xml":
            c = data.decode("utf-8")
            for tag in ("Application", "Company", "Template", "Manager"):
                c = re.sub(rf"<{tag}>.*?</{tag}>", f"<{tag}></{tag}>", c)
            data = c.encode("utf-8")
        elif re.search(r"word/(document|header\d*|footer\d*)\.xml$", name):
            c = data.decode("utf-8")
            c = re.sub(r"<w:t([^>]*)>([^<]*)</w:t>",
                       lambda m: f"<w:t{m.group(1)}></w:t>"
                       if any(s in m.group(2).lower() for s in _SYN) else m.group(0), c)
            data = c.encode("utf-8")
        # also fix the [Content_Types]/.rels reference to a dropped custom.xml
        if name == "[Content_Types].xml":
            c = data.decode("utf-8")
            c = c.replace('<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/>', "")
            data = c.encode("utf-8")
        if name == "_rels/.rels":
            c = data.decode("utf-8")
            c = re.sub(r'<Relationship[^>]*custom\.xml[^>]*/>', "", c)
            data = c.encode("utf-8")
        zout.writestr(item, data)
    zin.close(); zout.close(); os.replace(tmp, path)


_KM_IDS = ("korvin", "merrow", "km-6427819", "mercy vale", "csn-204418827",
           "fin-7740552", "vossmere", "travyn", "caldrane", "solthar", "halvek",
           "quenor", "riverbend", "5w-318", "5 west medic")


def verify_no_km_identifiers(path):
    """Hard gate: NO part may carry a Korvin Merrow world identifier (name, MRN, CSN,
    facility, or KM roster name). Catches header/footer leakage from the clone base
    that verify_no_synthetic does not. Raises naming the offending parts."""
    z = zipfile.ZipFile(path)
    offenders = []
    for n in z.namelist():
        if not (n.endswith(".xml") or n.endswith(".rels")):
            continue
        low = z.read(n).decode("utf-8", "ignore").lower()
        bad = [t for t in _KM_IDS if t in low]
        if bad:
            offenders.append((n, sorted(set(bad))))
    assert not offenders, f"KM identifier leak: {offenders}"
    print("  [OK ] verify_no_km_identifiers: zero KM identifiers in any part")
    return True


def verify_no_synthetic(path):
    """Hard gate: no part (every .xml and .rels, incl. ALL headers/footers and all
    three docProps) may contain a synthetic marker or authoring-tool fingerprint.
    Returns True or raises AssertionError naming the offending parts."""
    z = zipfile.ZipFile(path)
    offenders = []
    for n in z.namelist():
        if not (n.endswith(".xml") or n.endswith(".rels")):
            continue
        low = z.read(n).decode("utf-8", "ignore").lower()
        bad = [t for t in (_SYN + _TOOL) if t in low]
        if bad:
            offenders.append((n, sorted(set(bad))))
    assert not offenders, f"synthetic/tool token leak: {offenders}"
    print("  [OK ] verify_no_synthetic: zero synthetic/tool tokens in any part")
    return True


if __name__ == "__main__":
    print(__doc__)
    print("Import the primitives in a per-artifact build script. Example skeleton:")
    print('''
    from tools.mode_a_clone import *
    BASE = ".../platform/task2/current/discharge_summary_draft_incomplete_05242026.docx"
    OUT  = ".../task3/build/<new_task_file>.docx"
    d = clone(BASE, "/tmp/work.docx")
    paras = d.paragraphs
    set_text(paras[1], "<NEW TITLE>")
    set_text(paras[2], "Author: <name>  |  Department: <dept>  |  05/24/2026  |  Status: Signed")
    sig = paras[<n>]; delete_paragraphs(paras[<a>:<b>])
    for like, txt in <content>: insert_before(sig, like, txt)
    edit_band_cell(d, "Service", "Care Management", "Hospital Medicine")
    edit_footer(d, "Discharge Summary - Working Draft", "<new doc type>")
    d.save("/tmp/work.docx")
    integrity_gate("/tmp/work.docx")
    import shutil; shutil.copy("/tmp/work.docx", OUT)
    verify_against_base(OUT, BASE)
    ''')
