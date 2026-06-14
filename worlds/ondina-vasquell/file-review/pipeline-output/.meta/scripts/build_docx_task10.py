#!/usr/bin/env python3
"""Build the discharge instructions .docx file."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


OUTPUT = "/filesystem/tasks/task10_dc/started_discharge_instruction_draft_05212026.docx"

doc = Document()

# --- Page margins
for section in doc.sections:
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

# --- Base style
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

# Adjust heading styles
for level, size in [(1, 15), (2, 13), (3, 11)]:
    style = styles[f"Heading {level}"]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)


def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def add_horizontal_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "808080")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_para(text, bold=False, italic=False, size=11, align=None, space_before=0, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p


def add_mixed_para(segments, size=11, space_before=0, space_after=4, italic_all=False):
    """segments: list of (text, bold) tuples."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    for text, bold in segments:
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.bold = bold
        if italic_all:
            run.italic = True
    return p


def add_bullet(segments, size=11):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if isinstance(segments, str):
        segments = [(segments, False)]
    for text, bold in segments:
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.bold = bold
    return p


def add_hr_para():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    add_horizontal_rule(p)


def add_heading(text, level=1, space_before=10, space_after=4):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(space_before)
    h.paragraph_format.space_after = Pt(space_after)
    return h


# --- Header
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_after = Pt(2)
run = title_p.add_run("HARBOR CREST REGIONAL MEDICAL CENTER")
run.bold = True
run.font.size = Pt(15)
run.font.name = "Calibri"
run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

addr_p = doc.add_paragraph()
addr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
addr_p.paragraph_format.space_after = Pt(2)
arun = addr_p.add_run("2400 Mariners Bay Boulevard, Cypress Harbor, FL 33421  —  Main: (561) 555-0100")
arun.font.size = Pt(10.5)
arun.font.name = "Calibri"

add_hr_para()

main_title = doc.add_paragraph()
main_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
main_title.paragraph_format.space_before = Pt(2)
main_title.paragraph_format.space_after = Pt(10)
mt_run = main_title.add_run("DISCHARGE INSTRUCTIONS — DRAFT (NOT FINAL)")
mt_run.bold = True
mt_run.font.size = Pt(16)
mt_run.font.name = "Calibri"
mt_run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

# --- Patient info table
patient_rows = [
    ("Patient", "Ondina Vasquell", False),
    ("MRN", "OV-3358104", False),
    ("DOB", "03/14/1958", False),
    ("Unit / Room", "6 South Medicine / 6S-214", False),
    ("Draft date", "05/21/2026", False),
    ("Drafted by", "Hospital Medicine — Dr. Lillian Everet's team", False),
    ("Status", "STARTED DRAFT — to be completed at the discharge encounter", True),
    ("Preferred language",
     "Spanish (a certified Spanish interpreter will be used at delivery; the Spanish "
     "parallel section is to be completed by the translator before sign-out)",
     False),
]

table = doc.add_table(rows=len(patient_rows), cols=2)
table.autofit = False

for i, row_data in enumerate(patient_rows):
    label, value, bold_val = row_data
    row = table.rows[i]
    row.cells[0].width = Inches(1.8)
    row.cells[1].width = Inches(4.9)

    lc = row.cells[0]
    lc.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_shading(lc, "E8EEF5")
    lp = lc.paragraphs[0]
    lp.paragraph_format.space_after = Pt(0)
    lr = lp.add_run(label)
    lr.bold = True
    lr.font.size = Pt(10.5)
    lr.font.name = "Calibri"

    vc = row.cells[1]
    vc.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    vp = vc.paragraphs[0]
    vp.paragraph_format.space_after = Pt(0)
    vr = vp.add_run(value)
    vr.font.size = Pt(10.5)
    vr.font.name = "Calibri"
    if bold_val:
        vr.bold = True

tbl = table._tbl
tblPr = tbl.tblPr
tblBorders = OxmlElement("w:tblBorders")
for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
    b = OxmlElement(f"w:{border_name}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "4")
    b.set(qn("w:color"), "AAAAAA")
    tblBorders.append(b)
tblPr.append(tblBorders)

add_hr_para()

# --- Section 1
add_heading("1. Why You Were in the Hospital", level=1)
add_para(
    "You came to the hospital because you had a serious infection on the bottom of your "
    "left foot. The infection had spread into the soft tissue around the wound. You also "
    "had a flare-up of your kidney problem because of the infection. While you were here, "
    "the foot doctor (podiatrist) cleaned out the infected tissue, you received "
    "antibiotics through your IV, our team checked the blood flow to your foot, and we "
    "kept a close watch on your sugars and your kidneys. Your fever went away and your "
    "blood tests are getting better."
)
add_para(
    "This is still a serious foot wound. It needs careful follow-up at home or wherever "
    "you go next."
)
add_hr_para()

# --- Section 2
add_heading("2. Your Medicines at Home", level=1)
add_mixed_para([("Continue these medicines as you did before, with the same doses:", True)])

continue_meds = [
    [("Insulin glargine", True), (" 26 units under the skin, every night", False)],
    [("Insulin aspart", True), (" sliding scale under the skin, with meals (three times a day) — use the chart we give you", False)],
    [("Furosemide", True), (" 20 mg by mouth, once a day", False)],
    [("Atorvastatin", True), (" 40 mg by mouth, at bedtime", False)],
    [("Aspirin", True), (" 81 mg by mouth, once a day", False)],
    [("Clopidogrel", True), (" 75 mg by mouth, once a day", False)],
    [("Gabapentin", True), (" 300 mg by mouth, three times a day", False)],
    [("Ferrous sulfate", True), (" 325 mg by mouth, once a day", False)],
    [("Cholecalciferol (vitamin D)", True), (" 2,000 units by mouth, once a day", False)],
    [("Pantoprazole", True), (" 40 mg by mouth, once a day", False)],
    [("Acetaminophen (Tylenol)", True),
     (" 650 mg by mouth, up to three times a day ", False),
     ("only if needed", True),
     (" for your knee pain. Do NOT take ibuprofen, naproxen, or other anti-inflammatory pain pills — they are not safe for your kidneys.", False)],
    [("CPAP", True), (" — wear your CPAP mask every night when you sleep.", False)],
]
for med in continue_meds:
    add_bullet(med)

add_mixed_para(
    [("These medicines are On hold. Do NOT restart them on your own:", True)],
    space_before=6,
)
hold_meds = [
    [("metformin", True), (" — On hold — your hospital doctor will tell your regular doctor when to restart this.", False)],
    [("empagliflozin", True), (" — On hold — your hospital doctor will tell your regular doctor when to restart this.", False)],
    [("lisinopril", True), (" — On hold — your hospital doctor will tell your regular doctor when to restart this.", False)],
]
for med in hold_meds:
    add_bullet(med)

add_mixed_para(
    [("New antibiotic for your foot infection:", True)],
    space_before=6,
)
add_bullet("[DISCHARGE ANTIBIOTIC NAME, DOSE, DURATION — pending finalization]")
add_para(
    "Take this antibiotic exactly as written. Do not stop early, even if the foot looks better.",
    space_before=4,
)
add_hr_para()

# --- Section 3
add_heading("3. Wound Care", level=1)
wound_items = [
    [("A skilled nurse (home-health nurse or nurse at your next-care facility) will change the dressing on your left foot ", False),
     ("every day", True), (".", False)],
    [("The dressing is a ", False), ("moist dressing", True),
     (" — your nurse will keep the wound covered and slightly moist as instructed.", False)],
    [("Do NOT scrub", True),
     (" the wound. Do not pick at it. Do not put anything in it that the wound nurse has not given you.", False)],
    [("Keep the wound and the dressing ", False), ("clean and dry", True),
     (" between dressing changes. Do not soak the foot in water. Do not let the dressing get wet in the shower — your nurse will show you how to cover it.", False)],
    [("If the dressing falls off, comes loose, or gets soaked, call your home-health nurse or the wound-care number on this sheet.", False)],
]
for item in wound_items:
    add_bullet(item)
add_hr_para()

# --- Section 4
add_heading("4. Activity and Foot Offloading", level=1)

callout_table = doc.add_table(rows=1, cols=1)
callout_cell = callout_table.rows[0].cells[0]
set_cell_shading(callout_cell, "FFF4D6")
cp = callout_cell.paragraphs[0]
cp.paragraph_format.space_after = Pt(0)
cr = cp.add_run(
    "TO BE COMPLETED FROM CHART AT DISCHARGE ENCOUNTER — use PT/OT findings, "
    "wound plan, and case-management notes. Do not write generic clearance."
)
cr.bold = True
cr.font.size = Pt(10.5)
cr.font.name = "Calibri"
ctbl = callout_table._tbl
ctblPr = ctbl.tblPr
ctblBorders = OxmlElement("w:tblBorders")
for bn in ("top", "left", "bottom", "right"):
    b = OxmlElement(f"w:{bn}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:color"), "C9A227")
    ctblBorders.append(b)
ctblPr.append(ctblBorders)

add_para(
    "(This section is intentionally left blank in the draft. The discharging attending "
    "will complete it at the discharge encounter based on the documented physical therapy, "
    "occupational therapy, wound-care, and case-management findings in the chart. Do NOT "
    "pre-fill.)",
    italic=True,
    space_before=6,
)
add_hr_para()

# --- Section 5
add_heading("5. Follow-up Appointments", level=1)
add_para(
    "Please keep all of these appointments. Our discharge planner will help you schedule "
    "them before you leave."
)
followups = [
    [("Hospitalist or Primary Care Doctor — within 7 days.", True),
     (" Your primary care doctor is ", False),
     ("Esmeralda Torres-Hidalgo, MD", True),
     (", Bayfront Community Internal Medicine, 915 Pelican Drive, Cypress Harbor, FL 33421. Phone (561) 555-0823.", False)],
    [("Podiatry (foot doctor) — within 7 days.", True),
     (" With Dr. Priyanka Vell or her team, for wound check and continued foot care.", False)],
    [("Infectious Disease — within 1 week.", True),
     (" With Dr. Helena Brusk's clinic, to follow your antibiotic course.", False)],
    [("Vascular Surgery — referral to be sent (open).", True),
     (" The blood-flow study showed your blood flow to the foot needs to be looked at more closely. The hospital team will send the referral; the appointment date is still being arranged.", False)],
    [("Wound Care — within 3 to 5 days.", True),
     (" With the hospital wound-care clinic for measurement, photograph, and dressing-plan review.", False)],
]
for fu in followups:
    add_bullet(fu)
add_para(
    "If you have not heard about any of these appointments within 48 hours of getting "
    "home, please call us.",
    space_before=4,
)
add_hr_para()

# --- Section 6
add_heading("6. When to Call or Return to the Hospital", level=1)
add_mixed_para([
    ("Call your doctor or come back to the Emergency Department right away if you have any of these:", True)
])
warnings = [
    [("Fever of 38.0 °C (100.4 °F) or higher", True), (", chills, or shaking.", False)],
    [("New or worsening redness, warmth, swelling, or red streaks", True), (" around the foot wound.", False)],
    [("More drainage, pus, or a bad smell", True), (" coming from the wound.", False)],
    [("New pain in the foot or leg", True), (" — especially pain that wakes you up or is much worse than before.", False)],
    [("Signs of kidney problems:", True),
     (" less urine than usual, swelling in the legs that is new or worse, feeling very tired or confused, or nausea and vomiting that won't stop.", False)],
    [("Low blood sugar", True),
     (" (under 70 mg/dL): shaking, sweating, dizziness, confusion, or trouble waking up.", False)],
    [("High blood sugar", True),
     (" (over 300 mg/dL on two checks in a row), or feeling very thirsty, urinating a lot, vomiting, or feeling sick.", False)],
]
for w in warnings:
    add_bullet(w)
add_mixed_para([
    ("In an emergency — chest pain, trouble breathing, fainting, or a seizure — ", False),
    ("call 911", True),
    (".", False),
], space_before=6)
add_hr_para()

# --- Section 7
add_heading("7. Caregiver and Language Support", level=1)
caregiver_items = [
    [("A certified ", False), ("Spanish interpreter", True),
     (" will be used at the discharge encounter to review every section of these instructions with you. All teaching will be done in Spanish, with teach-back, before you sign out.", False)],
    [("Your daughter, ", False), ("Marisela", True),
     (", is invited and asked to be present during the discharge teaching. She will be included in the wound-care, medication, and foot-protection instructions so she can help support you at home.", False)],
    [("If you have questions after you go home and an interpreter is not available where you are, call the hospital main number — (561) 555-0100 — and ask for the Spanish interpreter line.", False)],
]
for c in caregiver_items:
    add_bullet(c)
add_hr_para()

# --- Section 8
add_heading("8. Spanish Parallel Section / Versión en Español", level=1)
add_para(
    "[To be completed by the certified Spanish translator before the discharge encounter. "
    "Each numbered section above (1–7) is to be rendered in Spanish in parallel, in the "
    "same order, using patient-facing plain language. Section 4 is to remain a placeholder "
    "in both languages until the discharging attending completes it from the chart.]",
    italic=True,
)

spanish_lines = [
    ("Sección 1 — Por qué estuvo en el hospital:", "[pendiente de traducción]"),
    ("Sección 2 — Sus medicamentos en casa:", "[pendiente de traducción]"),
    ("Sección 3 — Cuidado de la herida:", "[pendiente de traducción]"),
    ("Sección 4 — Actividad y descarga del pie:",
     "[A SER COMPLETADO DEL EXPEDIENTE EN EL MOMENTO DEL ALTA — no llenar con instrucciones genéricas.]"),
    ("Sección 5 — Citas de follow-up:", "[pendiente de traducción]"),
    ("Sección 6 — Cuándo llamar o regresar al hospital:", "[pendiente de traducción]"),
    ("Sección 7 — Apoyo familiar y de idioma — uso del interpreter:", "[pendiente de traducción]"),
]
for label, placeholder in spanish_lines:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r1 = p.add_run(label)
    r1.font.name = "Calibri"
    r1.font.size = Pt(11)
    r1.bold = True
    r2 = p.add_run(" " + placeholder)
    r2.font.name = "Calibri"
    r2.font.size = Pt(11)
    r2.italic = True

add_hr_para()

# --- Signature block
add_mixed_para([
    ("Discharging attending: ", False),
    ("Lillian Everet, MD", True),
    (" — to sign at discharge.", False),
], space_before=4)

add_para("Hospital Medicine, Harbor Crest Regional Medical Center", size=10.5)
add_para("NPI 1457629831  |  Pager (561) 555-0511", size=10.5, space_after=10)

sig_lines = [
    "Signature: ____________________________   Date: ____ / ____ / 2026   Time: ______",
    "Patient signature (acknowledging instructions received and reviewed with interpreter): ____________________________   Date: ____ / ____ / 2026",
    "Caregiver / daughter signature (if present): ____________________________   Date: ____ / ____ / 2026",
    "Interpreter (name and ID): ____________________________   Date: ____ / ____ / 2026",
]
for sline in sig_lines:
    add_para(sline, size=10.5, space_after=8)


doc.save(OUTPUT)
print(f"Wrote {OUTPUT}")

import os
size = os.path.getsize(OUTPUT)
print(f"Size: {size} bytes")
