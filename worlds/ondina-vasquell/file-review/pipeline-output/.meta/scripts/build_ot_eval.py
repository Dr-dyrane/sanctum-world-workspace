"""Build the OT evaluation .docx file."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


OUT_PATH = "/filesystem/therapy/ot_evaluation_05202026.docx"


def set_cell_shading(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def add_bottom_border(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "808080")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    add_bottom_border(p)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)


def styled_run(paragraph, text, *, bold=False, italic=False, size=10, color=None, font="Calibri"):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return run


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    styled_run(p, f"{label}: ", bold=True, size=10)
    styled_run(p, value, size=10)
    return p


def add_heading(doc, text, *, level=1):
    p = doc.add_paragraph()
    if level == 0:
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        styled_run(p, text, bold=True, size=14, color=RGBColor(0x1F, 0x3A, 0x5F))
    elif level == 1:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        styled_run(p, text, bold=True, size=12, color=RGBColor(0x1F, 0x3A, 0x5F))
        add_bottom_border(p)
    else:  # level 2
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        styled_run(p, text, bold=True, size=10.5, color=RGBColor(0x33, 0x33, 0x33))
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    styled_run(p, text, size=10)
    return p


def add_bullet_runs(doc, runs_spec):
    """runs_spec is a list of (text, bold) tuples."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    for text, bold in runs_spec:
        styled_run(p, text, bold=bold, size=10)
    return p


def add_numbered_runs(doc, runs_spec):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    for text, bold in runs_spec:
        styled_run(p, text, bold=bold, size=10)
    return p


def make_header_table(doc, rows):
    """rows: list of (header_list, value_list)."""
    table = doc.add_table(rows=len(rows) * 2, cols=len(rows[0][0]))
    table.autofit = True
    for i, (headers, values) in enumerate(rows):
        hrow = table.rows[i * 2]
        vrow = table.rows[i * 2 + 1]
        for j, h in enumerate(headers):
            cell = hrow.cells[j]
            cell.text = ""
            set_cell_shading(cell, "1F3A5F")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            styled_run(p, h, bold=True, size=9, color=RGBColor(0xFF, 0xFF, 0xFF))
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for j, v in enumerate(values):
            cell = vrow.cells[j]
            cell.text = ""
            set_cell_shading(cell, "F4F4F4")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            styled_run(p, v, size=10)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def main():
    doc = Document()

    # Default style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # --- Title block ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(0)
    styled_run(title, "Harbor Crest Regional Medical Center", bold=True, size=16,
               color=RGBColor(0x1F, 0x3A, 0x5F))

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(8)
    styled_run(sub, "Occupational Therapy — Confidential", italic=True, size=10,
               color=RGBColor(0x55, 0x55, 0x55))

    add_horizontal_rule(doc)

    # --- Header info tables ---
    make_header_table(
        doc,
        [
            (
                ["Patient", "Sex / DOB", "MRN"],
                ["Ondina Vasquell, 68 y", "Female / 03/14/1958", "OV-3358104"],
            ),
        ],
    )

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)

    make_header_table(
        doc,
        [
            (
                ["Date of Service", "Author", "Allergies", "Document"],
                [
                    "05/20/2026 1500",
                    "Sela Pruvost, OT (NPI 1771204938)",
                    "Sulfa (sulfonamide antibiotics) – rash",
                    "THERAPY EVALUATION",
                ],
            ),
        ],
    )

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    add_horizontal_rule(doc)

    # --- Main title ---
    add_heading(doc, "OCCUPATIONAL THERAPY INITIAL EVALUATION (EW14)", level=0)

    add_label_value(doc, "Author", "Sela Pruvost, OT — Occupational Therapy")
    add_label_value(doc, "NPI", "1771204938")
    add_label_value(doc, "Date of Service", "05/20/2026 1500")
    add_label_value(doc, "Status", "Signed")

    # --- PATIENT / ENCOUNTER ---
    add_heading(doc, "PATIENT / ENCOUNTER", level=1)
    pe = [
        ("Patient", "Ondina Vasquell, 68 y"),
        ("Sex / DOB", "Female / 03/14/1958"),
        ("MRN / FIN", "OV-3358104 / FIN-2207733"),
        ("Unit / Room", "6 South Medicine / 6S-214"),
        ("Code Status", "Full Code"),
        ("Allergies", "Sulfa (sulfonamide antibiotics) — rash"),
        ("Attending", "Lillian Everet, MD"),
        ("Service", "Hospital Medicine"),
        ("Date of Service", "05/20/2026"),
        ("Language", "Spanish preferred; certified hospital interpreter used for this session (Florencia Beltrán)"),
    ]
    for label, value in pe:
        add_bullet_runs(doc, [(f"{label}: ", True), (value, False)])

    # --- REASON FOR REFERRAL ---
    add_heading(doc, "REASON FOR REFERRAL", level=1)
    add_body(
        doc,
        "Initial occupational therapy evaluation for activities of daily living (ADL), self-care, and home-safety assessment in the context of a limb-threatening left diabetic foot infection with strict offloading precaution of the left forefoot. Referral specifically requests evaluation of the patient's ability to manage dressing, wound-protection, and offloading-device application at home given language preference, caregiver availability, and a second-floor walk-up residence.",
    )

    # --- DIAGNOSIS / CLINICAL CONTEXT ---
    add_heading(doc, "DIAGNOSIS / CLINICAL CONTEXT", level=1)
    add_body(
        doc,
        "Limb-threat left diabetic foot infection with offloading. Plantar left forefoot soft-tissue infection over the 1st–2nd metatarsal head, post bedside debridement by podiatry on 05/17/2026. Strict offloading of the left forefoot is in place per podiatry; weight-bearing of the left forefoot is restricted. Wound care is at a skilled-nursing level with daily moist wound dressings per CWOCN plan.",
    )
    add_body(
        doc,
        "Relevant comorbidities reviewed: insulin-dependent type 2 diabetes mellitus; diabetic peripheral neuropathy; CKD stage 3b; HFpEF; peripheral arterial disease (perfusion adequacy unresolved per vascular study); class I obesity; knee osteoarthritis; obstructive sleep apnea on CPAP; baseline limited mobility.",
    )

    # --- SUBJECTIVE ---
    add_heading(doc, "SUBJECTIVE", level=1)
    add_body(
        doc,
        "Patient was cooperative throughout the session. All exchanges occurred via the certified hospital Spanish interpreter; patient declined to proceed in English and her stated English proficiency is limited. Through the interpreter she stated, in substance, that she wants to be safe for home and that she is worried about getting back upstairs to her apartment. She acknowledged the wound on her left foot and stated that her daughter has been helping when she can. She reported no acute pain at rest while in bed and described the offloading boot as \"heavy and hot\" but indicated she understands she is supposed to wear it.",
    )
    add_body(
        doc,
        "She voiced that she would prefer not to go to a nursing facility and would rather go home. She did not voice a specific plan for who would help with her dressing or device application overnight or during her daughter's working hours.",
    )

    # --- OBJECTIVE ---
    add_heading(doc, "OBJECTIVE", level=1)

    add_heading(doc, "Self-care (ADL)", level=2)
    adl_items = [
        ("Upper-body dressing:", " Modified independent. Patient is able to don and doff a hospital gown and a loose shirt using bedside setup with minimal cueing."),
        ("Lower-body dressing:", " Limited by the left-foot offloading precaution. Patient is unable to manage pants, socks, and footwear over the offloading device without assistance."),
        ("Bathing / hygiene:", " Sponge bath at bedside with setup; lower-extremity hygiene requires assist secondary to the foot precaution and the device."),
        ("Toileting:", " Bedside commode with supervision per nursing; ambulation to bathroom is per PT plan with the front-wheeled walker."),
        ("Foot care / device adjustment:", " Patient cannot safely reach her left foot for dressing inspection, dressing protection, or independent adjustment of the offloading device. Reach is limited by trunk flexion tolerance, the bulk of the device itself, and the strict precaution against bearing weight or torque through the forefoot."),
    ]
    for label, rest in adl_items:
        add_bullet_runs(doc, [(label, True), (rest, False)])

    add_heading(doc, "Cognition / communication", level=2)
    cog_items = [
        "Alert and oriented; follows demonstrated tasks within session.",
        "Communication entirely interpreter-mediated this session; no English health-literacy material was used.",
        "No bilingual family member present at the bedside during the evaluation.",
    ]
    for it in cog_items:
        add_bullet_runs(doc, [(it, False)])

    add_heading(doc, "Vision / sensation", level=2)
    vs_items = [
        "Mild non-proliferative diabetic retinopathy bilaterally per outpatient ophthalmology; no acute visual complaint.",
        "Diabetic peripheral neuropathy of the feet documented; patient cannot reliably feel pressure or rubbing from the device against the wound site.",
    ]
    for it in vs_items:
        add_bullet_runs(doc, [(it, False)])

    add_heading(doc, "Home setup (per patient and daughter, interpreter-mediated)", level=2)
    home_items = [
        [("Lives alone in a second-floor walk-up apartment (1418 Calle del Mar, Apt 2B); no elevator; interior stairs to the unit.", False)],
        [("Adult daughter, Marisela Vasquell, is the primary caregiver. Marisela works night shifts as a CNA at a skilled-care facility; she is able to provide partial daytime support but is ", False),
         ("not", True),
         (" available overnight on workdays and is not available for daytime dressing changes on workdays.", False)],
        [("Spanish is the preferred and dominant home language; an interpreter is required for any teaching session in which the daughter is not present to translate.", False)],
        [("Existing home equipment: standard tub/shower, no grab bars, no raised toilet seat, no shower chair.", False)],
        [("No current home-health services in place.", False)],
    ]
    for spec in home_items:
        add_bullet_runs(doc, spec)

    add_heading(doc, "Offloading-device teach-back (this session)", level=2)
    teach_items = [
        [("Offloading-device application, dressing-protection technique, and the rationale for continuous wear during ambulation were taught with the interpreter at bedside.", False)],
        [("Teach-back was ", False), ("attempted", True), (" with the interpreter present.", False)],
        [("Patient was ", False), ("NOT yet able", True), (" to demonstrate correct offloading-device application or correct dressing protection independently when asked to show the steps back.", False)],
        [("Nursing offloading flowsheet for 05/19/2026 documents that the patient repeatedly removed the device at times during ambulation to the bathroom; the 05/20/2026 entry documents the device in place at rest but removed by the patient at times. This was reviewed with the patient through the interpreter.", False)],
    ]
    for spec in teach_items:
        add_bullet_runs(doc, spec)

    # --- ASSESSMENT ---
    add_heading(doc, "ASSESSMENT", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    styled_run(p, "Offloading and wound-protection ", size=10)
    styled_run(p, "teach-back NOT achieved", bold=True, size=10)
    styled_run(p, " this session.", size=10)

    add_body(
        doc,
        "The gap is attributable to identified system-level contributors that are present today and will remain present at the point of discharge unless mitigated:",
    )
    assessment_items = [
        [("Language access.", True), (" Teaching can only be delivered through a certified Spanish interpreter. The teaching that has occurred to date has been single-session, interpreter-mediated, without the patient holding written Spanish-language instructions at the bedside to rehearse between sessions. A single interpreter-mediated exposure to a multi-step motor task (device application, dressing protection, recognition of when to re-don the device) is not sufficient for return-demonstration in a patient with limited English proficiency and diabetic peripheral neuropathy of the feet.", False)],
        [("Caregiver schedule.", True), (" The primary caregiver, daughter Marisela, works night shifts and is not present at the bedside during routine daytime therapy or teaching windows. She has not yet been able to attend an interpreter-mediated teach-back session, and she is the person who would need to reinforce device use, dressing protection, and stair safety during the hours the patient is at home alone.", False)],
        [("Living environment.", True), (" Patient lives alone in a second-floor walk-up with interior stairs; the offloading device must be worn for any ambulation, including the trip to and from the apartment, and during any night-time bathroom trips when no caregiver is present.", False)],
        [("Sensory and ergonomic limits.", True), (" Diabetic peripheral neuropathy means the patient cannot rely on pain or pressure feedback to know when the device has shifted off the wound; class I obesity and trunk-flexion limits restrict her reach to her own left foot, which compounds the dressing-protection problem.", False)],
    ]
    for spec in assessment_items:
        add_numbered_runs(doc, spec)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    styled_run(p, "Home-safety risk for self-managed offloading and wound protection is therefore ", size=10)
    styled_run(p, "elevated", bold=True, size=10)
    styled_run(p, ", given solo living, stairs, partial caregiver availability, and language-access requirements for teaching. The failure to demonstrate the technique today reflects these contributors and the limited teaching exposure to date; it does not reflect unwillingness on the patient's part. Patient was engaged and cooperative throughout this session.", size=10)

    # --- PLAN ---
    add_heading(doc, "PLAN", level=1)
    plan_items = [
        [("Repeat offloading and dressing-protection teaching with the certified Spanish interpreter at each OT session.", True), (" Use return-demonstration as the standard for documenting whether teach-back is achieved; document each attempt and outcome in the flowsheet. Provide written Spanish-language step-by-step instructions for the patient to keep at the bedside between sessions.", False)],
        [("Coordinate caregiver training with the daughter (Marisela) when she is available.", True), (" Schedule at least one interpreter-mediated joint teach-back session that includes Marisela at the bedside, timed around her night-shift schedule. Document the joint session outcome separately from patient-only teach-back.", False)],
        [("Document home-safety contributors", True), (" (second-floor walk-up, lives alone, daughter on night shifts, Spanish-preferred and interpreter-required, diabetic peripheral neuropathy and reach limits) for the multidisciplinary team and for case management's disposition workup, so these contributors are visible to the discharge plan.", False)],
        [("Recommend continued skilled occupational therapy", True), (" during this admission, with reassessment of teach-back at each session. Coordinate with PT for stair assessment and with wound care/CWOCN for dressing-protection consistency.", False)],
        [("Coordinate with case management", True), (" on equipment needs (offloading device continuity at discharge, possible shower chair, raised toilet seat, grab-bar consideration) and on the level of post-acute support that would be required if teach-back continues to be unmet at the point of disposition.", False)],
        [("Coordinate with nursing and the interpreter line", True), (" so that every offloading-device application and removal at the bedside is paired with an interpreter-mediated cueing opportunity, not only formal therapy sessions.", False)],
    ]
    for spec in plan_items:
        add_numbered_runs(doc, spec)

    # --- GOALS ---
    add_heading(doc, "GOALS (initial set, to be reassessed each session)", level=1)
    goals_items = [
        [("Short-term (within next 2–3 sessions):", True), (" Patient will demonstrate correct offloading-device application with verbal cueing through the interpreter on at least two consecutive return-demonstrations.", False)],
        [("Short-term:", True), (" Patient will demonstrate dressing-protection technique (keeping the dressing dry and intact during dressing, transfers, and ambulation) with verbal cueing through the interpreter.", False)],
        [("Caregiver goal:", True), (" Daughter Marisela will participate in at least one interpreter-mediated teach-back session and will demonstrate correct device application as a reinforcer for the patient at home.", False)],
        [("Longer-term:", True), (" Patient and caregiver together will demonstrate consistent offloading-device use during all ambulation, including bathroom trips, without device removal.", False)],
    ]
    for spec in goals_items:
        add_bullet_runs(doc, spec)

    # --- DISPOSITION ---
    add_heading(doc, "DISPOSITION RECOMMENDATION (from OT perspective, contributing only)", level=1)
    add_body(
        doc,
        "OT does not make the disposition decision; the recommendation contributed from today's evaluation is that the home-safety picture is not yet established. Continued skilled therapy is recommended, and any home-with-services plan should be contingent on documented teach-back, caregiver participation, equipment in place, and a coordinated plan for the hours the patient is alone in the apartment.",
    )

    add_horizontal_rule(doc)

    sig = doc.add_paragraph()
    sig.paragraph_format.space_before = Pt(6)
    styled_run(
        sig,
        "Electronically signed by Sela Pruvost, OT (NPI 1771204938) on 05/20/2026 1500.",
        italic=True,
        size=9.5,
        color=RGBColor(0x55, 0x55, 0x55),
    )

    doc.save(OUT_PATH)
    print(f"Saved {OUT_PATH}")


if __name__ == "__main__":
    main()
